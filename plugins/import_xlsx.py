"""Импорт из XLSX-файла реестра."""
import os
import shutil
from pathlib import Path
from datetime import datetime
from openpyxl import load_workbook

__version__ = "1.0"
__description__ = "Импорт записей из XLSX-файла реестра"


def register(hooks, app, config, core):
    from flask import request, jsonify

    @app.route('/import/xlsx', methods=['POST'])
    def import_xlsx():
        if 'file' not in request.files:
            return jsonify({'success': False, 'error': 'Файл не загружен'}), 400
        f = request.files['file']
        if not f.filename or not f.filename.lower().endswith('.xlsx'):
            return jsonify({'success': False, 'error': 'Ожидается .xlsx'}), 400
        tmp = core.EXCEL_FILE + '.import_tmp'
        try:
            f.save(tmp)
        except Exception as e:
            return jsonify({'success': False, 'error': f'Ошибка сохранения: {e}'}), 500
        try:
            wb = load_workbook(tmp)
            ws = wb.active
            if ws.max_row < 1:
                raise ValueError('Пустой файл')
            headers = [c.value for c in ws[1]]
            if 'name_del' not in headers or 'link_del' not in headers:
                raise ValueError('Нет обязательных колонок')
        except Exception as e:
            try: os.remove(tmp)
            except Exception: pass
            return jsonify({'success': False, 'error': f'Ошибка чтения: {e}'}), 400
        try:
            if Path(core.EXCEL_FILE).exists():
                backup = str(Path(core.BASE_DIR) / f"реестр_{datetime.now().strftime('%Y%m%d_%H%M%S')}.bak")
                shutil.copy2(core.EXCEL_FILE, backup)
        except Exception as e:
            print(f"[IMPORT] Бэкап: {e}")
        try:
            shutil.move(tmp, core.EXCEL_FILE)
        except Exception as e:
            try: os.remove(tmp)
            except Exception: pass
            return jsonify({'success': False, 'error': f'Замена: {e}'}), 500
        try:
            wb = core.read_workbook()
            core.atomic_save_workbook(wb)
        except Exception as e:
            return jsonify({'success': False, 'error': f'Нормализация: {e}'}), 500
        core.invalidate_cache()
        core.save_json(str(Path(core.BASE_DIR) / 'card_order.json'), [])
        core.save_json(str(Path(core.BASE_DIR) / 'category_order.json'), [])
        core.log_action('import_xlsx', None, '', {'file': f.filename})
        core.run_hook('on_import', source='xlsx', added=0, extra={'file': f.filename})
        return jsonify({'success': True})

    hooks['ui_data_tab'].append({
        'id': 'import_xlsx_btn', 'icon': '📥',
        'label': 'Импорт из XLSX', 'action': 'openModal',
        'modal_id': 'importXlsxModal',
    })

    hooks['ui_modals'].append({
        'id': 'importXlsxModal', 'title': '📥 Импорт реестра из XLSX', 'width': 'normal',
        'html': '''
<div style="font-size:0.85rem; color:var(--meta-color); line-height:1.5; margin-bottom:14px;">
    <p>Файл заменит текущий реестр. Старый реестр будет сохранён как <b>.bak</b>-копия.</p>
    <p>Обязательные колонки: <code>name_del</code>, <code>link_del</code>.</p>
</div>
<label class="data-btn" for="pluginXlsxFile" style="justify-content:center;">
    <span class="data-icon">📁</span><span>Выбрать файл (.xlsx)</span>
</label>
<input type="file" id="pluginXlsxFile" accept=".xlsx" style="display:none;">
<div id="pluginXlsxProgress" style="display:none; margin-top:12px; color:var(--meta-color); font-size:0.9rem;"></div>
<div class="modal-actions">
    <button type="button" class="btn btn-secondary" onclick="closePluginModal('importXlsxModal')">Закрыть</button>
</div>
'''
    })

    hooks['ui_scripts'].append({'html': '''<script>
document.addEventListener('DOMContentLoaded', function() {
    const inp = document.getElementById('pluginXlsxFile');
    if (!inp) return;
    inp.addEventListener('change', function() {
        const file = this.files[0]; if (!file) return;
        if (!confirm('Импортировать ' + file.name + '? Текущий реестр заменится.')) { this.value = ''; return; }
        const prog = document.getElementById('pluginXlsxProgress');
        if (prog) { prog.style.display = 'block'; prog.textContent = '⏳ Загрузка...'; }
        const fd = new FormData(); fd.append('file', file);
        var _csrf = window.REESTR_CSRF || (document.querySelector('meta[name="csrf-token"]') || {getAttribute:function(){return ''}}).getAttribute('content') || '';
        fetch('/import/xlsx', {
            method: 'POST',
            body: fd
        })
            .then(r => r.json())
            .then(d => {
                if (d.success) {
                    if (prog) prog.textContent = '✅ Импортировано';
                    setTimeout(() => location.reload(), 800);
                } else { if (prog) prog.textContent = '❌ ' + (d.error || 'Ошибка'); }
            })
            .catch(() => { if (prog) prog.textContent = '❌ Ошибка сети'; })
            .finally(() => { inp.value = ''; });
    });
});
</script>'''})
