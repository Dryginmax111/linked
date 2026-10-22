"""
Абстракция хранилища: Excel (.xlsx) и SQLite (.db).

Позволяет работать в двух режимах без изменения остального кода:
- ExcelStorage  — текущее поведение (openpyxl + файл реестр.xlsx)
- SqliteStorage — SQLite с той же схемой в БД

Переключение — через settings.json → storage.type.
"""
import os
import re
import shutil
import sqlite3
import threading
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook

HEADERS = [
    'id', 'name_del', 'link_del', 'description', 'category',
    'priority', 'favorite', 'order_index', 'created_at', 'updated_at',
    'deleted_at', 'tags', 'click_count', 'last_clicked_at',
    'code', 'code_lang'
]

MAX_NAME_LEN = 200
MAX_DESC_LEN = 5000
MAX_TAGS_LEN = 500
MAX_CATEGORY_LEN = 200
MAX_URL_LEN = 4096

# Ограничение кода: SQLite легко хранит 5 МБ в TEXT,
# Excel не может больше 32 767 символов на ячейку.
MAX_CODE_LEN = 5_000_000        # для SQLite / JSON / API
EXCEL_CODE_LIMIT = 32_767       # жёсткий лимит Excel

ALLOWED_CODE_LANGS = {
    '', 'python', 'sql', 'javascript', 'typescript', 'html', 'css',
    'bash', 'shell', 'json', 'yaml', 'xml', 'text', 'plaintext',
    'php', 'ruby', 'go', 'rust', 'java', 'c', 'cpp', 'csharp',
    'powershell', 'ini', 'toml', 'markdown', 'diff',
}


def validate_code_lang(lang):
    lang = (lang or '').strip().lower()
    return lang if lang in ALLOWED_CODE_LANGS else ''


def safe_int(v, d=0):
    if v is None or v == '':
        return d
    try:
        return int(v)
    except (ValueError, TypeError):
        try:
            return int(float(v))
        except (ValueError, TypeError):
            return d


def safe_str(v, d=''):
    if v is None:
        return d
    return str(v)


def _truncate(v, n):
    if v is None:
        return ''
    s = str(v)
    return s if len(s) <= n else s[:n]


def now_ts():
    return datetime.now().strftime('%Y-%m-%d %H:%M:%S')


def is_valid_datetime_str(s):
    if not s:
        return False
    try:
        datetime.strptime(str(s), '%Y-%m-%d %H:%M:%S')
        return True
    except Exception:
        return False

def validate_sqlite_path(path_str, base_dir):
    """Проверяет путь для SQLite. Возвращает (ok, error, resolved_path, is_network).
    - Разрешает абсолютные пути и UNC.
    - Запрещает '..' в компонентах и системные каталоги.
    - Проверяет доступность родительской папки на запись.
    """
    s = (path_str or '').strip()
    if not s:
        return False, 'Пустой путь', None, False
    if '\x00' in s:
        return False, 'Недопустимые символы в пути', None, False

    is_net = _is_unc_path(s)

    # Нормализуем: для UNC не трогаем префикс, для остальных — Path
    if is_net:
        # \\server\share\dir\file.db
        # запрет '..' в компонентах
        parts = s.replace('/', '\\').split('\\')
        if any(p == '..' for p in parts):
            return False, 'Путь не должен содержать ".."', None, True
        resolved = s
    else:
        p = Path(s)
        if any(part == '..' for part in p.parts):
            return False, 'Путь не должен содержать ".."', None, False
        if not p.is_absolute():
            p = Path(base_dir) / p
        resolved = str(p)
        # запрет системных каталогов (грубая, но защищает от дурака)
        low = resolved.lower()
        forbidden = (
            'c:\\windows', 'c:\\program files', 'c:\\program files (x86)',
            '/etc/', '/usr/', '/var/', '/bin/', '/sbin/', '/boot/', '/sys/', '/proc/',
        )
        for f in forbidden:
            if low.startswith(f):
                return False, f'Запись в системный каталог запрещена: {f}', None, False

    # Родительская папка должна существовать и быть доступной на запись
    parent = Path(resolved).parent if not is_net else None
    if parent is not None:
        if not parent.exists():
            return False, f'Папка не существует: {parent}', None, False
        if not os.access(str(parent), os.W_OK):
            return False, f'Нет прав на запись в папку: {parent}', None, False
    # Если файл уже есть — проверим, что это валидный sqlite (первые 16 байт)
    if Path(resolved).exists():
        try:
            with open(resolved, 'rb') as f:
                head = f.read(16)
            if head and not head.startswith(b'SQLite format 3'):
                return False, f'Файл существует, но это не SQLite-база: {resolved}', None, is_net
        except Exception as e:
            return False, f'Не удалось прочитать файл: {e}', None, is_net
    # Touch-тест: пробуем создать временный файл рядом
    try:
        test_dir = Path(resolved).parent
        tmp = test_dir / f'.reestr_write_test_{os.getpid()}'
        tmp.write_bytes(b'x')
        tmp.unlink()
    except Exception as e:
        return False, f'Нет прав на запись: {e}', None, is_net

    return True, None, resolved, is_net
# ============================================================
# БАЗОВЫЙ КЛАСС
# ============================================================
class Storage:
    type = 'base'

    def read_workbook(self): raise NotImplementedError
    def atomic_save_workbook(self, wb): raise NotImplementedError
    def get_record_by_id(self, rid): raise NotImplementedError
    def add_record(self, name, link, desc='', category='', favorite=0, tags='',
                   code='', code_lang=''): raise NotImplementedError
    def update_record(self, rid, name, link, desc, category, favorite, tags='',
                      code=None, code_lang=None): raise NotImplementedError
    def soft_delete_record(self, rid): raise NotImplementedError
    def hard_delete_record(self, rid): raise NotImplementedError
    def restore_from_trash(self, rid): raise NotImplementedError
    def update_category_only(self, rid, cat): raise NotImplementedError
    def update_description(self, rid, desc): raise NotImplementedError
    def move_record_to_group(self, rid, group): raise NotImplementedError
    def toggle_favorite_record(self, rid): raise NotImplementedError
    def increment_click(self, rid): raise NotImplementedError
    def update_order(self, group_name, ordered_ids): raise NotImplementedError
    def sort_links_in_group(self, group_name, sort_by='name', asc=True): raise NotImplementedError
    def get_all_records(self): raise NotImplementedError
    def get_all_categories(self): raise NotImplementedError
    def get_all_tags(self): raise NotImplementedError
    def get_trash_count(self): raise NotImplementedError
    def get_next_id(self): raise NotImplementedError
    def import_from_xlsx(self, xlsx_path): raise NotImplementedError
    def close(self): pass


# ============================================================
# EXCEL STORAGE
# ============================================================
class ExcelStorage(Storage):
    type = 'excel'

    def __init__(self, excel_file, backup_file=None):
        self.excel_file = str(excel_file)
        self.backup_file = str(backup_file) if backup_file else self.excel_file + '.bak'
        self._lock = threading.RLock()

    def _is_locked_by_excel(self):
        try:
            base = Path(self.excel_file)
            lock = base.parent / ('~$' + base.name)
            return lock.exists()
        except Exception:
            return False

    def _safe_read(self):
        try:
            if Path(self.excel_file).exists():
                return load_workbook(self.excel_file)
        except Exception as e:
            print(f"[EXCEL] Ошибка чтения: {e}")
            if Path(self.backup_file).exists():
                try:
                    shutil.copy2(self.backup_file, self.excel_file)
                    return load_workbook(self.excel_file)
                except Exception:
                    pass
        wb = Workbook()
        ws = wb.active or wb.create_sheet("Sheet")
        ws.append(HEADERS)
        wb.save(self.excel_file)
        try:
            shutil.copy2(self.excel_file, self.backup_file)
        except Exception:
            pass
        return wb

    def _ensure_headers(self, wb):
        ws = wb.active or wb.create_sheet("Sheet")
        headers = [c.value for c in ws[1]] if ws.max_row >= 1 else []
        needs_save = False
        for col in HEADERS:
            if col not in headers:
                ws.cell(row=1, column=len(headers) + 1, value=col)
                headers.append(col)
                needs_save = True
        if needs_save:
            self.atomic_save_workbook(wb)
        return wb

    def read_workbook(self):
        return self._ensure_headers(self._safe_read())

    def atomic_save_workbook(self, wb):
        with self._lock:
            if self._is_locked_by_excel():
                raise RuntimeError(
                    "Файл реестр.xlsx сейчас открыт в Excel. Закройте его и попробуйте снова."
                )
            tmp = self.excel_file + '.tmp'
            try:
                wb.save(tmp)
                os.replace(tmp, self.excel_file)
            except Exception:
                try:
                    if Path(tmp).exists():
                        os.remove(tmp)
                except Exception:
                    pass
                raise
            try:
                shutil.copy2(self.excel_file, self.backup_file)
            except Exception:
                pass

    def _iter_rows(self, ws):
        for idx, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
            yield idx, row

    def get_next_id(self):
        wb = self.read_workbook(); ws = wb.active
        m = 0
        for row in ws.iter_rows(min_row=2, max_col=1, values_only=True):
            v = safe_int(row[0], 0)
            if v > m:
                m = v
        return m + 1

    def _next_order(self, ws, group_name):
        m = -1
        for row in ws.iter_rows(min_row=2, values_only=True):
            if row[1] == group_name:
                o = safe_int(row[7] if len(row) > 7 else 0, 0)
                if o > m:
                    m = o
        return m + 1

    def add_record(self, name, link, desc='', category='', favorite=0, tags='',
                   code='', code_lang=''):
        name = _truncate(name, MAX_NAME_LEN).strip()
        desc = _truncate(desc, MAX_DESC_LEN).strip()
        category = _truncate(category, MAX_CATEGORY_LEN).strip()
        tags = _truncate(tags, MAX_TAGS_LEN).strip()
        code = _truncate(code, EXCEL_CODE_LIMIT)
        code_lang = validate_code_lang(code_lang)
        if not name:
            return None
        # Разрешаем пустую ссылку, если есть код
        if not link and not code.strip():
            return None
        with self._lock:
            wb = self.read_workbook(); ws = wb.active
            new_id = self.get_next_id()
            order = self._next_order(ws, name)
            now = now_ts()
            ws.append([new_id, name, link, desc, category, '', favorite, order,
                       now, now, '', tags, 0, '', code, code_lang])
            self.atomic_save_workbook(wb)
        return new_id

    def update_record(self, rid, name, link, desc, category, favorite, tags='',
                      code=None, code_lang=None):
        name = _truncate(name, MAX_NAME_LEN).strip()
        desc = _truncate(desc, MAX_DESC_LEN).strip()
        category = _truncate(category, MAX_CATEGORY_LEN).strip()
        tags = _truncate(tags, MAX_TAGS_LEN).strip()
        with self._lock:
            wb = self.read_workbook(); ws = wb.active
            now = now_ts()
            for idx, row in self._iter_rows(ws):
                if safe_int(row[0], -1) == int(rid):
                    ws.cell(row=idx, column=2, value=name)
                    # ВСЕГДА обновляем ссылку (в т.ч. пустую) — согласовано с SqliteStorage
                    ws.cell(row=idx, column=3, value=link or '')
                    ws.cell(row=idx, column=4, value=desc)
                    ws.cell(row=idx, column=5, value=category)
                    ws.cell(row=idx, column=7, value=favorite)
                    ws.cell(row=idx, column=10, value=now)
                    ws.cell(row=idx, column=12, value=tags)
                    if code is not None:
                        ws.cell(row=idx, column=15, value=_truncate(code, EXCEL_CODE_LIMIT))
                    if code_lang is not None:
                        ws.cell(row=idx, column=16, value=validate_code_lang(code_lang))
                    break
            self.atomic_save_workbook(wb)

    def soft_delete_record(self, rid):
        with self._lock:
            wb = self.read_workbook(); ws = wb.active
            now = now_ts()
            for idx, row in self._iter_rows(ws):
                if safe_int(row[0], -1) == int(rid):
                    ws.cell(row=idx, column=11, value=now)
                    break
            self.atomic_save_workbook(wb)

    def hard_delete_record(self, rid):
        name = None
        with self._lock:
            wb = self.read_workbook(); ws = wb.active
            for idx, row in self._iter_rows(ws):
                if safe_int(row[0], -1) == int(rid):
                    name = safe_str(row[1])
                    ws.delete_rows(idx)
                    break
            self.atomic_save_workbook(wb)
        return name

    def restore_from_trash(self, rid):
        changed = False
        with self._lock:
            wb = self.read_workbook(); ws = wb.active
            for idx, row in self._iter_rows(ws):
                if safe_int(row[0], -1) == int(rid):
                    if not safe_str(row[10] if len(row) > 10 else ''):
                        return False
                    ws.cell(row=idx, column=11, value='')
                    changed = True
                    break
            if changed:
                self.atomic_save_workbook(wb)
        return changed

    def update_category_only(self, rid, cat):
        cat = _truncate(cat, MAX_CATEGORY_LEN).strip()
        old = ''
        name = ''
        with self._lock:
            wb = self.read_workbook(); ws = wb.active
            now = now_ts()
            found = False
            for idx, row in self._iter_rows(ws):
                if safe_int(row[0], -1) == int(rid):
                    old = safe_str(row[4])
                    name = safe_str(row[1])
                    ws.cell(row=idx, column=5, value=cat)
                    ws.cell(row=idx, column=10, value=now)
                    found = True
                    break
            if not found:
                return None
            self.atomic_save_workbook(wb)
        return {'old': old, 'name': name, 'new': cat}

    def update_description(self, rid, desc):
        desc = _truncate(desc, MAX_DESC_LEN).strip()
        with self._lock:
            wb = self.read_workbook(); ws = wb.active
            now = now_ts()
            found = False
            for idx, row in self._iter_rows(ws):
                if safe_int(row[0], -1) == int(rid):
                    ws.cell(row=idx, column=4, value=desc)
                    ws.cell(row=idx, column=10, value=now)
                    found = True
                    break
            if not found:
                return False
            self.atomic_save_workbook(wb)
        return True

    def move_record_to_group(self, rid, group):
        group = _truncate(group, MAX_NAME_LEN).strip()
        old = ''
        with self._lock:
            wb = self.read_workbook(); ws = wb.active
            now = now_ts()
            new_order = self._next_order(ws, group)
            found = False
            for idx, row in self._iter_rows(ws):
                if safe_int(row[0], -1) == int(rid):
                    old = safe_str(row[1])
                    ws.cell(row=idx, column=2, value=group)
                    ws.cell(row=idx, column=8, value=new_order)
                    ws.cell(row=idx, column=10, value=now)
                    found = True
                    break
            if not found:
                return None
            self.atomic_save_workbook(wb)
        return {'old': old, 'new': group}

    def toggle_favorite_record(self, rid):
        with self._lock:
            wb = self.read_workbook(); ws = wb.active
            for idx, row in self._iter_rows(ws):
                if safe_int(row[0], -1) == int(rid):
                    cur = safe_int(row[6] if len(row) > 6 else 0, 0)
                    ws.cell(row=idx, column=7, value=1 if cur == 0 else 0)
                    break
            self.atomic_save_workbook(wb)

    def increment_click(self, rid):
        result = None
        with self._lock:
            wb = self.read_workbook(); ws = wb.active
            now = now_ts()
            for idx, row in self._iter_rows(ws):
                if safe_int(row[0], -1) == int(rid):
                    cur = safe_int(row[12] if len(row) > 12 else 0, 0)
                    cnt = cur + 1
                    ws.cell(row=idx, column=13, value=cnt)
                    ws.cell(row=idx, column=14, value=now)
                    result = {'count': cnt, 'name': safe_str(row[1]), 'link': safe_str(row[2])}
                    break
            self.atomic_save_workbook(wb)
        return result

    def update_order(self, group_name, ordered_ids):
        with self._lock:
            wb = self.read_workbook(); ws = wb.active
            id_to_row = {}
            for idx, row in self._iter_rows(ws):
                if row[1] == group_name:
                    id_to_row[safe_int(row[0], -1)] = idx
            for order, rid in enumerate(ordered_ids):
                rid = safe_int(rid, -1)
                if rid in id_to_row:
                    ws.cell(row=id_to_row[rid], column=8, value=order)
            self.atomic_save_workbook(wb)

    def sort_links_in_group(self, group_name, sort_by='name', asc=True):
        with self._lock:
            wb = self.read_workbook(); ws = wb.active
            rows = []
            for idx, row in self._iter_rows(ws):
                if row[1] == group_name:
                    rows.append({'idx': idx, 'row': row})
            if not rows:
                return False
            if sort_by == 'name':
                rows.sort(key=lambda x: safe_str(x['row'][2]).lower(), reverse=not asc)
            elif sort_by == 'date':
                rows.sort(key=lambda x: safe_str(x['row'][9]), reverse=not asc)
            elif sort_by == 'type':
                def tk(x):
                    link = safe_str(x['row'][2])
                    if link.startswith('\\\\') or link.startswith('//'):
                        return '3_net'
                    if '://' in link:
                        from urllib.parse import urlparse
                        host = urlparse(link).netloc.split(':')[0]
                        if host.startswith('192.168.') or host.startswith('10.'):
                            return '1_lan'
                        return '0_wan'
                    return '2_other'
                rows.sort(key=tk, reverse=not asc)
            for o, it in enumerate(rows):
                ws.cell(row=it['idx'], column=8, value=o)
            self.atomic_save_workbook(wb)
        return True

    def get_record_by_id(self, rid):
        wb = self.read_workbook(); ws = wb.active
        for row in ws.iter_rows(min_row=2, values_only=True):
            if safe_int(row[0], -1) == int(rid):
                return {
                    'id': row[0], 'name_del': row[1], 'link_del': row[2],
                    'description': row[3] if len(row) > 3 else '',
                    'category': row[4] if len(row) > 4 else '',
                    'favorite': safe_int(row[6] if len(row) > 6 else 0, 0),
                    'order_index': safe_int(row[7] if len(row) > 7 else 0, 0),
                    'created_at': row[8] if len(row) > 8 else '',
                    'updated_at': row[9] if len(row) > 9 else '',
                    'deleted_at': row[10] if len(row) > 10 else '',
                    'tags': row[11] if len(row) > 11 else '',
                    'click_count': safe_int(row[12] if len(row) > 12 else 0, 0),
                    'last_clicked_at': row[13] if len(row) > 13 else '',
                    'code': safe_str(row[14]) if len(row) > 14 else '',
                    'code_lang': safe_str(row[15]) if len(row) > 15 else '',
                }
        return None
    
    def get_all_records(self):
        wb = self.read_workbook(); ws = wb.active
        out = []
        for row in ws.iter_rows(min_row=2, values_only=True):
            if not row or row[0] is None:
                continue
            out.append({
                'id': safe_int(row[0], 0),
                'name_del': safe_str(row[1] if len(row) > 1 else ''),
                'link_del': safe_str(row[2] if len(row) > 2 else ''),
                'description': safe_str(row[3] if len(row) > 3 else ''),
                'category': safe_str(row[4] if len(row) > 4 else ''),
                'priority': safe_str(row[5] if len(row) > 5 else ''),
                'favorite': safe_int(row[6] if len(row) > 6 else 0, 0),
                'order_index': safe_int(row[7] if len(row) > 7 else 0, 0),
                'created_at': safe_str(row[8] if len(row) > 8 else ''),
                'updated_at': safe_str(row[9] if len(row) > 9 else ''),
                'deleted_at': safe_str(row[10] if len(row) > 10 else ''),
                'tags': safe_str(row[11] if len(row) > 11 else ''),
                'click_count': safe_int(row[12] if len(row) > 12 else 0, 0),
                'last_clicked_at': safe_str(row[13] if len(row) > 13 else ''),
                'code': safe_str(row[14]) if len(row) > 14 else '',
                'code_lang': safe_str(row[15]) if len(row) > 15 else '',
            })
        return out

    def get_all_categories(self):
        cats = set()
        for r in self.get_all_records():
            if r['deleted_at']:
                continue
            c = r['category'].strip()
            if c:
                cats.add(c)
        return sorted(cats)

    def get_all_tags(self):
        tags = {}
        for r in self.get_all_records():
            if r['deleted_at']:
                continue
            s = r['tags']
            if s:
                for t in s.split(','):
                    t = t.strip()
                    if t:
                        tags[t] = tags.get(t, 0) + 1
        return tags

    def get_trash_count(self):
        c = 0
        for r in self.get_all_records():
            if r['deleted_at'] and is_valid_datetime_str(r['deleted_at']):
                c += 1
        return c

    def import_from_xlsx(self, xlsx_path):
        if not Path(xlsx_path).exists():
            return 0
        src = load_workbook(xlsx_path)
        ws = src.active
        headers = [c.value for c in ws[1]]
        idx = {h: i for i, h in enumerate(headers) if h}
        rows = list(ws.iter_rows(min_row=2, values_only=True))

        wb = Workbook(); ws2 = wb.active or wb.create_sheet("Sheet")
        ws2.append(HEADERS)
        for row in rows:
            def pick(name, default=''):
                i = idx.get(name)
                if i is None or i >= len(row):
                    return default
                return row[i] if row[i] is not None else default
            ws2.append([
                safe_int(pick('id'), 0),
                safe_str(pick('name_del')),
                safe_str(pick('link_del')),
                safe_str(pick('description')),
                safe_str(pick('category')),
                safe_str(pick('priority')),
                safe_int(pick('favorite'), 0),
                safe_int(pick('order_index'), 0),
                safe_str(pick('created_at')),
                safe_str(pick('updated_at')),
                safe_str(pick('deleted_at')),
                safe_str(pick('tags')),
                safe_int(pick('click_count'), 0),
                safe_str(pick('last_clicked_at')),
                safe_str(pick('code')),
                safe_str(pick('code_lang')),
            ])
        with self._lock:
            self.atomic_save_workbook(wb)
        return len(rows)


# ============================================================
# SQLITE STORAGE
# ============================================================
_SQLITE_SCHEMA = """
CREATE TABLE IF NOT EXISTS links (
    id              INTEGER PRIMARY KEY,
    name_del        TEXT NOT NULL,
    link_del        TEXT DEFAULT '',
    description     TEXT DEFAULT '',
    category        TEXT DEFAULT '',
    priority        TEXT DEFAULT '',
    favorite        INTEGER DEFAULT 0,
    order_index     INTEGER DEFAULT 0,
    created_at      TEXT DEFAULT '',
    updated_at      TEXT DEFAULT '',
    deleted_at      TEXT DEFAULT '',
    tags            TEXT DEFAULT '',
    click_count     INTEGER DEFAULT 0,
    last_clicked_at TEXT DEFAULT '',
    code            TEXT DEFAULT '',
    code_lang       TEXT DEFAULT ''
);
CREATE INDEX IF NOT EXISTS idx_links_group ON links(name_del);
CREATE INDEX IF NOT EXISTS idx_links_category ON links(category);
CREATE INDEX IF NOT EXISTS idx_links_deleted ON links(deleted_at);
CREATE INDEX IF NOT EXISTS idx_links_favorite ON links(favorite);
"""


def _migrate_sqlite(conn):
    try:
        cur = conn.execute('PRAGMA table_info(links)')
        cols = {r[1] for r in cur.fetchall()}
        if 'code' not in cols:
            conn.execute("ALTER TABLE links ADD COLUMN code TEXT DEFAULT ''")
        if 'code_lang' not in cols:
            conn.execute("ALTER TABLE links ADD COLUMN code_lang TEXT DEFAULT ''")
        conn.commit()
    except Exception as e:
        print(f"[SQLITE-MIGRATE] {e}")

def _is_unc_path(p):
    """True, если путь похож на UNC (\\\\server\\share) или сетевой."""
    s = str(p)
    return s.startswith('\\\\') or s.startswith('//')

class SqliteStorage(Storage):
    type = 'sqlite'

    def __init__(self, db_path, network_safe=False):
        self.db_path = str(db_path)
        self._lock = threading.RLock()
        self._network = bool(network_safe) or _is_unc_path(self.db_path)

        # Для сетевых путей — обычный журнал (WAL на SMB ломается)
        # Для локальных — WAL (быстрее).
        try:
            self._conn = sqlite3.connect(self.db_path, check_same_thread=False, timeout=30)
        except sqlite3.Error as e:
            raise RuntimeError(
                f"Не удалось открыть SQLite-базу «{self.db_path}»: {e}. "
                f"Проверьте, что путь доступен на запись."
            ) from e
        self._conn.row_factory = sqlite3.Row
        try:
            if self._network:
                self._conn.execute('PRAGMA journal_mode=DELETE')
                self._conn.execute('PRAGMA synchronous=FULL')
            else:
                self._conn.execute('PRAGMA journal_mode=WAL')
                self._conn.execute('PRAGMA synchronous=NORMAL')
            self._conn.execute('PRAGMA busy_timeout=30000')
            self._conn.execute('PRAGMA foreign_keys=ON')
        except Exception:
            pass
        with self._lock:
            self._conn.executescript(_SQLITE_SCHEMA)
            self._conn.commit()
            _migrate_sqlite(self._conn)

    def is_network(self):
        return self._network

    def close(self):
        try:
            self._conn.close()
        except Exception:
            pass

    def _row_to_dict(self, row):
        d = {
            'id': row['id'],
            'name_del': row['name_del'] or '',
            'link_del': row['link_del'] or '',
            'description': row['description'] or '',
            'category': row['category'] or '',
            'priority': row['priority'] or '',
            'favorite': int(row['favorite'] or 0),
            'order_index': int(row['order_index'] or 0),
            'created_at': row['created_at'] or '',
            'updated_at': row['updated_at'] or '',
            'deleted_at': row['deleted_at'] or '',
            'tags': row['tags'] or '',
            'click_count': int(row['click_count'] or 0),
            'last_clicked_at': row['last_clicked_at'] or '',
        }
        try:
            d['code'] = row['code'] or ''
            d['code_lang'] = row['code_lang'] or ''
        except (IndexError, KeyError):
            d['code'] = ''
            d['code_lang'] = ''
        return d

    def get_next_id(self):
        with self._lock:
            cur = self._conn.execute('SELECT COALESCE(MAX(id), 0) + 1 AS nid FROM links')
            return int(cur.fetchone()['nid'])

    def _next_order(self, group_name):
        cur = self._conn.execute(
            'SELECT COALESCE(MAX(order_index), -1) + 1 AS n FROM links WHERE name_del = ?',
            (group_name,)
        )
        return int(cur.fetchone()['n'])

    def add_record(self, name, link, desc='', category='', favorite=0, tags='',
                   code='', code_lang=''):
        name = _truncate(name, MAX_NAME_LEN).strip()
        desc = _truncate(desc, MAX_DESC_LEN).strip()
        category = _truncate(category, MAX_CATEGORY_LEN).strip()
        tags = _truncate(tags, MAX_TAGS_LEN).strip()
        code = _truncate(code, MAX_CODE_LEN)
        code_lang = validate_code_lang(code_lang)
        if not name:
            return None
        if not link and not code.strip():
            return None
        with self._lock:
            new_id = self.get_next_id()
            order = self._next_order(name)
            now = now_ts()
            self._conn.execute(
                '''INSERT INTO links
                   (id, name_del, link_del, description, category, priority,
                    favorite, order_index, created_at, updated_at, deleted_at,
                    tags, click_count, last_clicked_at, code, code_lang)
                   VALUES (?, ?, ?, ?, ?, '', ?, ?, ?, ?, '', ?, 0, '', ?, ?)''',
                (new_id, name, link, desc, category, int(favorite), order,
                 now, now, tags, code, code_lang)
            )
            self._conn.commit()
        return new_id

    def update_record(self, rid, name, link, desc, category, favorite, tags='',
                      code=None, code_lang=None):
        name = _truncate(name, MAX_NAME_LEN).strip()
        desc = _truncate(desc, MAX_DESC_LEN).strip()
        category = _truncate(category, MAX_CATEGORY_LEN).strip()
        tags = _truncate(tags, MAX_TAGS_LEN).strip()
        with self._lock:
            now = now_ts()
            sets = ['name_del=?', 'description=?', 'category=?',
                    'favorite=?', 'updated_at=?', 'tags=?']
            params = [name, desc, category, int(favorite), now, tags]
            # Всегда обновляем link_del, даже если он пустой —
            # это позволяет очистить ссылку у code-only записи
            sets.insert(1, 'link_del=?')
            params.insert(1, link or '')
            if code is not None:
                sets.append('code=?')
                params.append(_truncate(code, MAX_CODE_LEN))
            if code_lang is not None:
                sets.append('code_lang=?')
                params.append(validate_code_lang(code_lang))
            params.append(int(rid))
            self._conn.execute(
                'UPDATE links SET ' + ', '.join(sets) + ' WHERE id=?',
                tuple(params)
            )
            self._conn.commit()

    def soft_delete_record(self, rid):
        with self._lock:
            self._conn.execute(
                'UPDATE links SET deleted_at = ? WHERE id = ?',
                (now_ts(), int(rid))
            )
            self._conn.commit()

    def hard_delete_record(self, rid):
        with self._lock:
            cur = self._conn.execute('SELECT name_del FROM links WHERE id = ?', (int(rid),))
            row = cur.fetchone()
            name = row['name_del'] if row else None
            self._conn.execute('DELETE FROM links WHERE id = ?', (int(rid),))
            self._conn.commit()
        return name

    def restore_from_trash(self, rid):
        with self._lock:
            cur = self._conn.execute(
                'SELECT deleted_at FROM links WHERE id = ?', (int(rid),)
            )
            row = cur.fetchone()
            if not row or not row['deleted_at']:
                return False
            self._conn.execute(
                'UPDATE links SET deleted_at = "" WHERE id = ?', (int(rid),)
            )
            self._conn.commit()
        return True

    def update_category_only(self, rid, cat):
        cat = _truncate(cat, MAX_CATEGORY_LEN).strip()
        with self._lock:
            cur = self._conn.execute(
                'SELECT category, name_del FROM links WHERE id = ?', (int(rid),)
            )
            row = cur.fetchone()
            if not row:
                return None
            old = row['category'] or ''
            name = row['name_del'] or ''
            self._conn.execute(
                'UPDATE links SET category = ?, updated_at = ? WHERE id = ?',
                (cat, now_ts(), int(rid))
            )
            self._conn.commit()
        return {'old': old, 'name': name, 'new': cat}

    def update_description(self, rid, desc):
        desc = _truncate(desc, MAX_DESC_LEN).strip()
        with self._lock:
            cur = self._conn.execute(
                'UPDATE links SET description = ?, updated_at = ? WHERE id = ?',
                (desc, now_ts(), int(rid))
            )
            self._conn.commit()
        return cur.rowcount > 0

    def move_record_to_group(self, rid, group):
        group = _truncate(group, MAX_NAME_LEN).strip()
        with self._lock:
            cur = self._conn.execute(
                'SELECT name_del FROM links WHERE id = ?', (int(rid),)
            )
            row = cur.fetchone()
            if not row:
                return None
            old = row['name_del'] or ''
            order = self._next_order(group)
            self._conn.execute(
                'UPDATE links SET name_del = ?, order_index = ?, updated_at = ? WHERE id = ?',
                (group, order, now_ts(), int(rid))
            )
            self._conn.commit()
        return {'old': old, 'new': group}

    def toggle_favorite_record(self, rid):
        with self._lock:
            self._conn.execute(
                'UPDATE links SET favorite = CASE WHEN favorite = 0 THEN 1 ELSE 0 END WHERE id = ?',
                (int(rid),)
            )
            self._conn.commit()

    def increment_click(self, rid):
        with self._lock:
            cur = self._conn.execute(
                'SELECT name_del, link_del, click_count FROM links WHERE id = ?', (int(rid),)
            )
            row = cur.fetchone()
            if not row:
                return None
            cnt = int(row['click_count'] or 0) + 1
            self._conn.execute(
                'UPDATE links SET click_count = ?, last_clicked_at = ? WHERE id = ?',
                (cnt, now_ts(), int(rid))
            )
            self._conn.commit()
        return {'count': cnt, 'name': row['name_del'] or '', 'link': row['link_del'] or ''}

    def update_order(self, group_name, ordered_ids):
        with self._lock:
            for order, rid in enumerate(ordered_ids):
                try:
                    self._conn.execute(
                        'UPDATE links SET order_index = ? WHERE id = ? AND name_del = ?',
                        (order, int(rid), group_name)
                    )
                except Exception:
                    pass
            self._conn.commit()

    def sort_links_in_group(self, group_name, sort_by='name', asc=True):
        order_dir = 'ASC' if asc else 'DESC'
        if sort_by == 'name':
            col = 'name_del'
        elif sort_by == 'date':
            col = 'updated_at'
        elif sort_by == 'type':
            col = 'link_del'
        else:
            col = 'order_index'
        with self._lock:
            cur = self._conn.execute(
                f'SELECT id FROM links WHERE name_del = ? ORDER BY {col} {order_dir}',
                (group_name,)
            )
            ids = [r['id'] for r in cur.fetchall()]
            if not ids:
                return False
            for o, rid in enumerate(ids):
                self._conn.execute(
                    'UPDATE links SET order_index = ? WHERE id = ?', (o, rid)
                )
            self._conn.commit()
        return True

    def get_record_by_id(self, rid):
        with self._lock:
            cur = self._conn.execute('SELECT * FROM links WHERE id = ?', (int(rid),))
            row = cur.fetchone()
        return self._row_to_dict(row) if row else None

    def get_all_records(self):
        with self._lock:
            cur = self._conn.execute('SELECT * FROM links ORDER BY id')
            return [self._row_to_dict(r) for r in cur.fetchall()]

    def get_all_categories(self):
        with self._lock:
            cur = self._conn.execute(
                '''SELECT DISTINCT category FROM links
                   WHERE deleted_at = '' AND category != ''
                   ORDER BY category COLLATE NOCASE'''
            )
            return [r['category'] for r in cur.fetchall()]

    def get_all_tags(self):
        tags = {}
        with self._lock:
            cur = self._conn.execute(
                "SELECT tags FROM links WHERE deleted_at = '' AND tags != ''"
            )
            for r in cur.fetchall():
                for t in (r['tags'] or '').split(','):
                    t = t.strip()
                    if t:
                        tags[t] = tags.get(t, 0) + 1
        return tags

    def get_trash_count(self):
        with self._lock:
            cur = self._conn.execute(
                "SELECT COUNT(*) AS c FROM links WHERE deleted_at != ''"
            )
            return int(cur.fetchone()['c'])

    def read_workbook(self):
        wb = Workbook(); ws = wb.active or wb.create_sheet("Sheet")
        ws.append(HEADERS)
        for r in self.get_all_records():
            ws.append([
                r['id'], r['name_del'], r['link_del'], r['description'],
                r['category'], r['priority'], r['favorite'], r['order_index'],
                r['created_at'], r['updated_at'], r['deleted_at'],
                r['tags'], r['click_count'], r['last_clicked_at'],
                r.get('code', ''), r.get('code_lang', ''),
            ])
        return wb

    def atomic_save_workbook(self, wb):
        ws = wb.active
        with self._lock:
            self._conn.execute('BEGIN IMMEDIATE')
            try:
                self._conn.execute('DELETE FROM links')
                for row in ws.iter_rows(min_row=2, values_only=True):
                    if not row or row[0] is None:
                        continue
                    self._conn.execute(
                        '''INSERT INTO links
                           (id, name_del, link_del, description, category, priority,
                            favorite, order_index, created_at, updated_at, deleted_at,
                            tags, click_count, last_clicked_at, code, code_lang)
                           VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)''',
                        (
                            safe_int(row[0], 0),
                            safe_str(row[1] if len(row) > 1 else ''),
                            safe_str(row[2] if len(row) > 2 else ''),
                            safe_str(row[3] if len(row) > 3 else ''),
                            safe_str(row[4] if len(row) > 4 else ''),
                            safe_str(row[5] if len(row) > 5 else ''),
                            safe_int(row[6] if len(row) > 6 else 0, 0),
                            safe_int(row[7] if len(row) > 7 else 0, 0),
                            safe_str(row[8] if len(row) > 8 else ''),
                            safe_str(row[9] if len(row) > 9 else ''),
                            safe_str(row[10] if len(row) > 10 else ''),
                            safe_str(row[11] if len(row) > 11 else ''),
                            safe_int(row[12] if len(row) > 12 else 0, 0),
                            safe_str(row[13] if len(row) > 13 else ''),
                            safe_str(row[14] if len(row) > 14 else ''),
                            safe_str(row[15] if len(row) > 15 else ''),
                        )
                    )
                self._conn.commit()
            except Exception:
                self._conn.rollback()
                raise

    def import_from_xlsx(self, xlsx_path):
        if not Path(xlsx_path).exists():
            return 0
        src = load_workbook(xlsx_path); ws = src.active
        headers = [c.value for c in ws[1]]
        idx = {h: i for i, h in enumerate(headers) if h}
        rows = list(ws.iter_rows(min_row=2, values_only=True))

        def pick(row, name, default=''):
            i = idx.get(name)
            if i is None or i >= len(row):
                return default
            return row[i] if row[i] is not None else default

        inserted = 0
        with self._lock:
            self._conn.execute('BEGIN IMMEDIATE')
            try:
                for row in rows:
                    rid = safe_int(pick(row, 'id'), 0)
                    if not rid:
                        continue
                    try:
                        self._conn.execute(
                            '''INSERT OR REPLACE INTO links
                               (id, name_del, link_del, description, category, priority,
                                favorite, order_index, created_at, updated_at, deleted_at,
                                tags, click_count, last_clicked_at, code, code_lang)
                               VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)''',
                            (
                                rid,
                                safe_str(pick(row, 'name_del')),
                                safe_str(pick(row, 'link_del')),
                                safe_str(pick(row, 'description')),
                                safe_str(pick(row, 'category')),
                                safe_str(pick(row, 'priority')),
                                safe_int(pick(row, 'favorite'), 0),
                                safe_int(pick(row, 'order_index'), 0),
                                safe_str(pick(row, 'created_at')),
                                safe_str(pick(row, 'updated_at')),
                                safe_str(pick(row, 'deleted_at')),
                                safe_str(pick(row, 'tags')),
                                safe_int(pick(row, 'click_count'), 0),
                                safe_str(pick(row, 'last_clicked_at')),
                                safe_str(pick(row, 'code')),
                                safe_str(pick(row, 'code_lang')),
                            )
                        )
                        inserted += 1
                    except Exception as e:
                        print(f"[SQLITE-IMPORT] строка {rid}: {e}")
                self._conn.commit()
            except Exception:
                self._conn.rollback()
                raise
        return inserted


# ============================================================
# ФАБРИКА
# ============================================================
def get_storage(settings, base_dir):
    base_dir = Path(base_dir)
    cfg = settings.get('storage') or {}
    st_type = (cfg.get('type') or 'sqlite').lower()

    if st_type == 'sqlite':
        sqlite_cfg = cfg.get('sqlite') or {}
        db_rel = sqlite_cfg.get('path') or 'reestr.db'
        db_path = Path(db_rel)
        if not db_path.is_absolute() and not _is_unc_path(str(db_path)):
            db_path = base_dir / db_rel
        db_str = str(db_path)

        excel_file = str(base_dir / 'реестр.xlsx')
        first_run = not Path(db_str).exists()

        try:
            storage = SqliteStorage(db_str)
        except RuntimeError:
            raise
        except Exception as e:
            raise RuntimeError(f"Не удалось инициализировать SQLite: {e}") from e

        if first_run and Path(excel_file).exists():
            try:
                n = storage.import_from_xlsx(excel_file)
                print(f"[STORAGE] Импортировано из xlsx: {n} записей")
            except Exception as e:
                print(f"[STORAGE] Ошибка импорта xlsx: {e}")
        else:
            try:
                total = len(storage.get_all_records())
                if total == 0 and Path(excel_file).exists():
                    n = storage.import_from_xlsx(excel_file)
                    print(f"[STORAGE] Пустая БД — импортировано {n} записей из xlsx")
            except Exception:
                pass
        return storage

    excel_file = str(base_dir / 'реестр.xlsx')
    excel_backup = str(base_dir / 'реестр.bak')
    storage = ExcelStorage(excel_file, excel_backup)
    if not Path(excel_file).exists():
        wb = Workbook(); ws = wb.active or wb.create_sheet("Sheet")
        ws.append(HEADERS)
        storage.atomic_save_workbook(wb)
    return storage
