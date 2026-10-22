from flask import Flask, request, redirect, render_template, jsonify, send_file, g, has_request_context, session
from openpyxl import Workbook, load_workbook
from pathlib import Path
from datetime import datetime, timedelta
from markupsafe import escape
import atexit
import os
import re
import json
import time
import shutil
import secrets
import threading
import logging
import logging.handlers
import signal
import sys

from plugins_loader import (
    load_plugins, run_hook, get_loaded_plugins, HOOK_NAMES,
    get_plugins_info, read_plugin_code, save_plugin_code,
    delete_plugin, set_enabled, validate_plugin_name, check_syntax,
    collect_ui_slots, get_custom_themes, add_custom_theme, delete_custom_theme,
    get_all_hooks, reset_hooks, reload_plugins,
)

from storage import (
    get_storage, ExcelStorage, SqliteStorage, Storage,
    HEADERS as STORAGE_HEADERS,
    safe_int as _st_safe_int,
    safe_str as _st_safe_str,
    validate_code_lang,
    validate_sqlite_path,          # NEW
)

# ============================================================
# ПУТИ И КОНСТАНТЫ
# ============================================================
BASE_DIR = Path(__file__).resolve().parent
EXCEL_FILE = str(BASE_DIR / "реестр.xlsx")
EXCEL_BACKUP = str(BASE_DIR / "реестр.bak")
CATEGORY_COLORS_FILE = str(BASE_DIR / "category_colors.json")
CARD_ORDER_FILE = str(BASE_DIR / "card_order.json")
CATEGORY_ORDER_FILE = str(BASE_DIR / "category_order.json")
GROUP_ICONS_FILE = str(BASE_DIR / "group_icons.json")
SETTINGS_FILE = str(BASE_DIR / "settings.json")
LOGS_FILE = str(BASE_DIR / "logs.json")
PLUGINS_DIR = str(BASE_DIR / "plugins")
PLUGINS_CONFIG_FILE = str(BASE_DIR / "plugins.json")
CUSTOM_THEMES_FILE = str(BASE_DIR / "custom_themes.json")
BACKUP_DIR = str(BASE_DIR / "backups")
SECRET_KEY_FILE = str(BASE_DIR / ".flask_secret")
SCHEMA_VERSION_FILE = str(BASE_DIR / "schema_version.json")

BACKUP_KEEP_DAYS = 30
TRASH_KEEP_DAYS = 30
SCHEMA_VERSION = 1
MAX_NAME_LEN, MAX_DESC_LEN, MAX_TAGS_LEN, MAX_CATEGORY_LEN, MAX_URL_LEN = 200, 5000, 500, 200, 4096
MAX_CODE_LEN = 20000

HEADERS = [
    'id', 'name_del', 'link_del', 'description', 'category',
    'priority', 'favorite', 'order_index', 'created_at', 'updated_at',
    'deleted_at', 'tags', 'click_count', 'last_clicked_at',
    'code', 'code_lang'
]


# ============================================================
# ЛОГИРОВАНИЕ
# ============================================================
def _setup_logging():
    log_dir = BASE_DIR / 'logs'
    log_dir.mkdir(exist_ok=True)
    log_file = log_dir / 'app.log'

    handler = logging.handlers.RotatingFileHandler(
        str(log_file), maxBytes=5 * 1024 * 1024, backupCount=5, encoding='utf-8'
    )
    handler.setFormatter(logging.Formatter(
        '%(asctime)s [%(levelname)-7s] %(name)s | %(message)s',
        datefmt='%Y-%m-%d %H:%M:%S'
    ))

    console = logging.StreamHandler(sys.stdout)
    console.setFormatter(logging.Formatter(
        '%(asctime)s [%(levelname)-7s] %(message)s',
        datefmt='%H:%M:%S'
    ))

    root = logging.getLogger()
    root.setLevel(logging.INFO)
    root.handlers.clear()
    root.addHandler(handler)
    root.addHandler(console)

    logging.getLogger('werkzeug').setLevel(logging.WARNING)


_setup_logging()
log = logging.getLogger('reestr')


# ============================================================
# SECRET KEY
# ============================================================
def _load_secret_key():
    try:
        if Path(SECRET_KEY_FILE).exists():
            key = Path(SECRET_KEY_FILE).read_text(encoding='utf-8').strip()
            if len(key) >= 32:
                return key
        new_key = secrets.token_hex(32)
        Path(SECRET_KEY_FILE).write_text(new_key, encoding='utf-8')
        try:
            os.chmod(SECRET_KEY_FILE, 0o600)
        except Exception:
            pass
        log.info("Сгенерирован новый secret_key: %s", Path(SECRET_KEY_FILE).name)
        return new_key
    except Exception as e:
        log.error("Не удалось сохранить secret_key: %s", e)
        return secrets.token_hex(32)


# ============================================================
# FLASK APP
# ============================================================
app = Flask(__name__, static_folder=str(BASE_DIR / 'static'), template_folder=str(BASE_DIR / 'templates'))
app.secret_key = _load_secret_key()
app.config['SESSION_COOKIE_SAMESITE'] = 'Lax'
app.config['SESSION_COOKIE_HTTPONLY'] = True
try:
    app.json.ensure_ascii = False
except Exception:
    app.config['JSON_AS_ASCII'] = False

ALLOWED_ORIGINS = ['http://127.0.0.1:5432', 'http://localhost:5432', 'http://[::1]:5432']
_extra = os.environ.get('REESTR_ALLOWED_ORIGINS', '')
if _extra:
    for _o in _extra.split(','):
        _o = _o.strip()
        if _o and _o not in ALLOWED_ORIGINS:
            ALLOWED_ORIGINS.append(_o)
ALLOWED_ORIGINS = tuple(ALLOWED_ORIGINS)

CSRF_EXEMPT_PATHS = {'/api/csrf-token', '/api/js-log'}
CSRF_EXEMPT_METHODS = {'GET', 'HEAD', 'OPTIONS'}


# ============================================================
# CSRF
# ============================================================
def _get_csrf_token():
    if 'csrf_token' not in session:
        session['csrf_token'] = secrets.token_urlsafe(32)
    return session['csrf_token']


def _check_csrf():
    if request.method in CSRF_EXEMPT_METHODS:
        return None
    if request.path in CSRF_EXEMPT_PATHS:
        return None

    sent = request.headers.get('X-CSRF-Token', '')
    if not sent:
        sent = (request.form.get('_csrf_token') or '')
    if not sent:
        try:
            j = request.get_json(silent=True) or {}
            sent = j.get('_csrf_token', '')
        except Exception:
            sent = ''

    expected = session.get('csrf_token', '')
    has_cookie = bool(request.cookies.get('session'))

    ok = bool(expected) and bool(sent) and secrets.compare_digest(sent, expected)
    if not ok:
        log.warning(
            "[CSRF] Заблокирован: %s %s | sent=%s expected=%s cookie=%s origin=%s",
            request.method, request.path,
            (sent[:8] + '…') if sent else '<empty>',
            (expected[:8] + '…') if expected else '<empty>',
            'yes' if has_cookie else 'no',
            request.headers.get('Origin', '-') or '-'
        )
        return jsonify({
            'success': False,
            'error': 'CSRF: недействительный токен',
            'code': 'CSRF'
        }), 403
    return None


@app.context_processor
def _inject_csrf():
    return {'csrf_token': _get_csrf_token()}


@app.route('/api/csrf-token')
def api_csrf_token():
    return jsonify({'token': _get_csrf_token()})


@app.route('/api/csrf-debug')
def api_csrf_debug():
    t = _get_csrf_token()
    return jsonify({
        'session_has_token': bool(session.get('csrf_token')),
        'session_token_prefix': (t[:8] + '…') if t else '',
        'has_cookie': bool(request.cookies.get('session')),
        'cookie_keys': list(request.cookies.keys()),
        'host': request.host,
        'host_url': request.host_url,
        'scheme': request.scheme,
    })


@app.route('/api/js-log', methods=['POST'])
def api_js_log():
    """Принимает логи и ошибки JavaScript с клиента.
    Пишет в консоль сервера с префиксом [JS] или [JS-ERROR]."""
    try:
        raw = request.get_data(as_text=True) or ''
        try:
            data = json.loads(raw)
        except Exception:
            data = {'message': raw[:500]}

        kind = str(data.get('kind', 'log'))[:20]
        msg = str(data.get('message', ''))[:2000]
        extra = str(data.get('extra', ''))[:2000]

        if kind == 'error':
            log.error("[JS-ERROR] %s", msg)
            if extra:
                for line in extra.split('\n')[:10]:
                    log.error("[JS-ERROR]   %s", line)
        else:
            log.info("[JS] %s", msg)
            if extra:
                for line in extra.split('\n')[:10]:
                    log.info("[JS]   %s", line)
    except Exception as e:
        log.error("[JS-LOG] Не удалось обработать: %s", e)
    return '', 204


@app.before_request
def _check_origin_and_csrf():
    if request.method in ('POST', 'PUT', 'DELETE', 'PATCH'):
        origin = request.headers.get('Origin', '')
        if origin:
            if origin in ALLOWED_ORIGINS:
                pass
            else:
                try:
                    if origin.rstrip('/') == request.host_url.rstrip('/'):
                        pass
                    else:
                        log.warning("[CSRF] Заблокирован Origin=%s", origin)
                        return jsonify({'success': False, 'error': 'CSRF: origin', 'code': 'CSRF'}), 403
                except Exception:
                    return jsonify({'success': False, 'error': 'CSRF: origin', 'code': 'CSRF'}), 403
        csrf_resp = _check_csrf()
        if csrf_resp is not None:
            return csrf_resp


# ============================================================
# STORAGE
# ============================================================
_storage = None


def get_current_storage():
    global _storage
    if _storage is None:
        _storage = get_storage(load_settings(), BASE_DIR)
        log.info("[STORAGE] Тип: %s", _storage.type)
    return _storage


def switch_storage(new_type, extra=None):
    global _storage
    cur = load_settings()
    storage_cfg = cur.get('storage') or {}
    storage_cfg['type'] = new_type
    # NEW: помечаем, что пользователь выбрал вручную —
    # это отключает принудительную миграцию Excel → SQLite при рестарте.
    storage_cfg['_explicit'] = True
    if extra:
        storage_cfg.update(extra)
    cur['storage'] = storage_cfg
    save_settings(cur)
    if _storage is not None:
        try:
            _storage.close()
        except Exception:
            pass
    _storage = get_storage(cur, BASE_DIR)
    invalidate_cache(full=True)
    log.info("[STORAGE] Переключено на: %s", _storage.type)
    return _storage


def _auto_migrate_to_sqlite(base_dir):
    """
    Одноразовая миграция Excel → SQLite при старте сервера.

    Условия срабатывания:
      - файл-маркер .sqlite_migrated отсутствует
      - в settings.json storage.type != "sqlite"
      - НЕТ флага _explicit (пользователь явно не выбирал Excel)
      - есть что импортировать (реестр.xlsx), либо БД уже пуста/отсутствует

    Логика:
      - если в reestr.db уже есть данные — просто переключаем тип в settings
      - если БД пуста или отсутствует, а xlsx есть — импортируем
      - если xlsx нет — просто переключаем тип (пустая БД)
      - если импорт упал — не переключаем (чтобы не потерять данные)

    После успеха создаётся маркер .sqlite_migrated — миграция больше не сработает.
    """
    marker = Path(base_dir) / '.sqlite_migrated'
    if marker.exists():
        return

    settings = _load_json(SETTINGS_FILE, {})
    if not isinstance(settings, dict):
        settings = {}
    cfg = settings.get('storage') or {}
    cur_type = str(cfg.get('type') or 'excel').lower()

    # Уже на SQLite — просто ставим маркер и выходим
    if cur_type == 'sqlite':
        marker.write_text('already-sqlite', encoding='utf-8')
        return

    # NEW: пользователь явно выбрал Excel — не мигрируем принудительно.
    if cfg.get('_explicit') and cur_type == 'excel':
        marker.write_text('user-chose-excel', encoding='utf-8')
        log.info("[MIGRATION] Пользователь явно выбрал Excel — миграция пропущена")
        return

    excel_file = Path(base_dir) / 'реестр.xlsx'
    db_file = Path(base_dir) / 'reestr.db'

    # Проверяем, есть ли данные в существующей БД
    db_has_data = False
    if db_file.exists():
        try:
            st = SqliteStorage(str(db_file))
            try:
                db_has_data = len(st.get_all_records()) > 0
            finally:
                st.close()
        except Exception as e:
            log.error("[MIGRATION] Проверка БД: %s", e)

    imported = 0
    if not db_has_data and excel_file.exists():
        log.info("[MIGRATION] Импорт %s → %s", excel_file.name, db_file.name)
        try:
            st = SqliteStorage(str(db_file))
            try:
                imported = st.import_from_xlsx(str(excel_file))
            finally:
                st.close()
        except Exception as e:
            log.error("[MIGRATION] Импорт не удался: %s", e)
            return
    elif not excel_file.exists():
        log.info("[MIGRATION] Файл %s не найден — переключаюсь на пустую SQLite",
                 excel_file.name)

    # Обновляем settings.json
    settings.setdefault('storage', {})
    settings['storage']['type'] = 'sqlite'
    settings['storage'].setdefault('sqlite', {'path': 'reestr.db'})
    _save_json(SETTINGS_FILE, settings)

    marker.write_text(f'migrated:imported={imported}', encoding='utf-8')
    log.info("[MIGRATION] Переключено на SQLite (импортировано: %d записей)", imported)


# ============================================================
# КЭШ
# ============================================================
_excel_write_lock = threading.RLock()
_CACHE = {
    'groups': None, 'categories': None,
    'category_counts': None, 'total': 0, 'all_names': None,
    'all_tags': None, 'trash_count': 0, 'uncategorized_count': 0,
    'url_index': None,
}


def invalidate_cache(full=True, groups=False, categories=False, tags=False,
                     url_index=False, trash=False):
    global _CACHE
    if full:
        _CACHE = {
            'groups': None, 'categories': None,
            'category_counts': None, 'total': 0, 'all_names': None,
            'all_tags': None, 'trash_count': 0, 'uncategorized_count': 0,
            'url_index': None,
        }
        return
    if groups: _CACHE['groups'] = None
    if categories:
        _CACHE['categories'] = None
        _CACHE['category_counts'] = None
    if tags: _CACHE['all_tags'] = None
    if url_index: _CACHE['url_index'] = None
    if trash: _CACHE['trash_count'] = None


# ============================================================
# УТИЛИТЫ
# ============================================================
_FORBIDDEN_URL_SCHEMES = ('javascript:', 'data:', 'vbscript:')


def sanitize_link(link):
    if not link:
        return ''
    link = str(link).strip()
    if len(link) > MAX_URL_LEN:
        return ''
    head = re.sub(r'[\s\x00-\x1f]', '', link[:64]).lower()
    for bad in _FORBIDDEN_URL_SCHEMES:
        if head.startswith(bad):
            return ''
    return link


def _truncate(v, n):
    if v is None:
        return ''
    s = str(v)
    return s if len(s) <= n else s[:n]


def normalize_url_for_dup(url):
    if not url:
        return ''
    u = url.strip().lower()
    u = re.sub(r'^https?://', '', u)
    u = re.sub(r'^file://', '', u)
    u = u.rstrip('/')
    u = re.sub(r'^www\.', '', u)
    return u.split('?')[0].split('#')[0]


_HEX_COLOR_RE = re.compile(r'^#[0-9a-fA-F]{3}([0-9a-fA-F]{3})?([0-9a-fA-F]{2})?$')


def is_valid_hex_color(c):
    return bool(c and _HEX_COLOR_RE.match(c))


def _parse_ids(ids):
    if not isinstance(ids, list):
        return None
    try:
        return set(int(x) for x in ids if x is not None)
    except (ValueError, TypeError):
        return None


def _sanitize_toolbar_btn(btn):
    if not isinstance(btn, dict):
        return None
    url = str(btn.get('url', ''))[:500].strip()
    if url and not (url.startswith('http://') or url.startswith('https://')
                    or url.startswith('#category:')):
        url = ''
    return {'icon': str(btn.get('icon', '🔗'))[:8],
            'label': str(btn.get('label', ''))[:50], 'url': url}


def safe_int(v, d=0):
    if v is None or v == '':
        return d
    try:
        return int(v)
    except (ValueError, TypeError):
        try:
            return int(float(v))
        except (ValueError, TypeError):
            return d


def safe_str(v, d=''):
    if v is None:
        return d
    return str(v)


def is_valid_datetime_str(s):
    if not s:
        return False
    try:
        datetime.strptime(str(s), '%Y-%m-%d %H:%M:%S')
        return True
    except Exception:
        return False


# ============================================================
# SCHEMA VERSION
# ============================================================
def _load_schema_version():
    try:
        if Path(SCHEMA_VERSION_FILE).exists():
            d = json.loads(Path(SCHEMA_VERSION_FILE).read_text(encoding='utf-8'))
            if isinstance(d, dict):
                return int(d.get('version', 0))
    except Exception:
        pass
    return 0


def _save_schema_version(v):
    try:
        Path(SCHEMA_VERSION_FILE).write_text(
            json.dumps({'version': v, 'updated_at': datetime.now().strftime('%Y-%m-%d %H:%M:%S')},
                       ensure_ascii=False, indent=2),
            encoding='utf-8'
        )
    except Exception as e:
        log.error("Не удалось сохранить schema_version: %s", e)


# ============================================================
# ОБЁРТКИ STORAGE
# ============================================================
def safe_read_workbook():
    return get_current_storage().read_workbook()


def read_workbook():
    if has_request_context():
        cached = getattr(g, '_reestr_wb', None)
        if cached is not None:
            return cached
    wb = get_current_storage().read_workbook()
    if has_request_context():
        g._reestr_wb = wb
    return wb


def atomic_save_workbook(wb):
    get_current_storage().atomic_save_workbook(wb)
    # NEW: сбрасываем per-request кэш книги, чтобы последующие
    # read_workbook() в этом же запросе увидели свежие данные.
    if has_request_context():
        try:
            g._reestr_wb = None
        except Exception:
            pass


def _ensure_daily_backup():
    Path(BACKUP_DIR).mkdir(parents=True, exist_ok=True)
    today = datetime.now().strftime('%Y-%m-%d')
    target = Path(BACKUP_DIR) / f"реестр_{today}.xlsx"
    if target.exists():
        return
    try:
        shutil.copy2(EXCEL_FILE, target)
    except Exception:
        return
    try:
        cutoff = time.time() - BACKUP_KEEP_DAYS * 86400
        for f in Path(BACKUP_DIR).glob("реестр_*.xlsx"):
            try:
                if f.stat().st_mtime < cutoff:
                    f.unlink()
            except Exception:
                pass
    except Exception:
        pass


def get_next_id(ws=None):
    return get_current_storage().get_next_id()


def get_next_order_index(ws, name_del):
    st = get_current_storage()
    if st.type == 'sqlite':
        return st._next_order(name_del)
    return st._next_order(ws, name_del)


# ============================================================
# JSON
# ============================================================
def _try_recover_json(path):
    try:
        content = open(path, 'r', encoding='utf-8').read().strip()
    except Exception:
        return None
    if not content:
        return None
    items = []
    for line in content.split('\n'):
        line = line.strip().rstrip(',')
        if not line or line in ('[', ']', '{', '}'):
            continue
        try:
            items.append(json.loads(line))
        except json.JSONDecodeError:
            pass
    if items:
        if all(isinstance(x, list) for x in items):
            flat = []
            for x in items:
                flat.extend(x)
            return flat
        if len(items) == 1:
            return items[0]
        return items
    if '][' in content:
        parts, depth, start, i = [], 0, None, 0
        while i < len(content):
            c = content[i]
            if c in '([{':
                if depth == 0:
                    start = i
                depth += 1
            elif c in ')]}':
                depth -= 1
                if depth == 0 and start is not None:
                    try:
                        parts.append(json.loads(content[start:i + 1]))
                    except json.JSONDecodeError:
                        pass
                    start = None
            i += 1
        if parts:
            if all(isinstance(x, list) for x in parts):
                flat = []
                for x in parts:
                    flat.extend(x)
                return flat
            return parts[0] if len(parts) == 1 else parts
    for trim in range(len(content) - 1, 0, -1):
        if content[trim] in ']}':
            try:
                return json.loads(content[:trim + 1])
            except json.JSONDecodeError:
                continue
    return None


def _load_json(path, default):
    if not Path(path).exists():
        return default
    try:
        with open(path, 'r', encoding='utf-8') as f:
            return json.load(f)
    except json.JSONDecodeError as e:
        log.warning("[JSON] %s повреждён: %s", path, e)
        recovered = _try_recover_json(path)
        if recovered is not None:
            try:
                bad = path + f'.bad-{datetime.now().strftime("%Y%m%d_%H%M%S")}'
                shutil.copy2(path, bad)
                log.info("[JSON] Битый сохранён: %s", Path(bad).name)
            except Exception:
                pass
            try:
                with open(path, 'w', encoding='utf-8') as f:
                    json.dump(recovered, f, ensure_ascii=False, indent=2)
                log.info("[JSON] %s восстановлен", Path(path).name)
            except Exception:
                pass
            return recovered
        return default
    except Exception as e:
        log.error("[JSON] Ошибка %s: %s", path, e)
        return default


def _save_json(path, data):
    try:
        with open(path, 'w', encoding='utf-8') as f:
            json.dump(data, f, ensure_ascii=False, indent=2)
    except Exception as e:
        log.error("[JSON] Ошибка записи %s: %s", path, e)


def load_category_colors(): return _load_json(CATEGORY_COLORS_FILE, {})
def save_category_colors(c): _save_json(CATEGORY_COLORS_FILE, c)


def load_card_order():
    d = _load_json(CARD_ORDER_FILE, [])
    return d if isinstance(d, list) else []


def save_card_order(o): _save_json(CARD_ORDER_FILE, o)


def load_category_order():
    d = _load_json(CATEGORY_ORDER_FILE, [])
    return d if isinstance(d, list) else []


def save_category_order(o): _save_json(CATEGORY_ORDER_FILE, o)


def load_group_icons():
    d = _load_json(GROUP_ICONS_FILE, {})
    return d if isinstance(d, dict) else {}


def save_group_icons(i): _save_json(GROUP_ICONS_FILE, i)


def load_settings():
    d = _load_json(SETTINGS_FILE, {})
    if not isinstance(d, dict):
        d = {}
    # NEW: тема по умолчанию — «Светлая»
    d.setdefault('theme', 'light')
    d.setdefault('custom_toolbar_buttons', [])
    d.setdefault('profile', {'name': 'Admin', 'color': '#f5d76e'})
    d.setdefault('shimmer_color_1', '')
    d.setdefault('shimmer_color_2', '')
    d.setdefault('header_shimmer', 'sweep')
    d.setdefault('card_shimmer', True)
    d.setdefault('shimmer_speed', 'normal')
    d.setdefault('shimmer_intensity', 1.0)
    d.setdefault('header_anim', True)
    d.setdefault('mesh_style', 'off')
    d.setdefault('card_strip', 'left')
    d.setdefault('storage', {'type': 'sqlite', 'sqlite': {'path': 'reestr.db'}})
    return d


def save_settings(s): _save_json(SETTINGS_FILE, s)


# ============================================================
# ЛОГИ ДЕЙСТВИЙ
# ============================================================
def log_action(action, target_id=None, target_name='', details=None):
    try:
        logs = _load_json(LOGS_FILE, [])
        if not isinstance(logs, list):
            logs = []
        logs = [l for l in logs if isinstance(l, dict)]
        logs.append({
            'ts': datetime.now().strftime('%Y-%m-%d %H:%M:%S'),
            'action': action,
            'id': target_id,
            'name': str(target_name)[:120] if target_name else '',
            'details': details or {}
        })
        cutoff = (datetime.now() - timedelta(days=180)).strftime('%Y-%m-%d %H:%M:%S')
        logs = [l for l in logs if str(l.get('ts', '')) >= cutoff][-500:]
        _save_json(LOGS_FILE, logs)
    except Exception as e:
        log.error("[LOG] %s", e)


# ============================================================
# CRUD
# ============================================================
def add_record(name, link, desc='', category='', favorite=0, tags='',
               code='', code_lang=''):
    """Добавить запись. Разрешает пустую ссылку, если есть код."""
    name = _truncate(name, MAX_NAME_LEN).strip()
    desc = _truncate(desc, MAX_DESC_LEN).strip()
    category = _truncate(category, MAX_CATEGORY_LEN).strip()
    tags = _truncate(tags, MAX_TAGS_LEN).strip()
    code = _truncate(code, MAX_CODE_LEN)
    code_lang = validate_code_lang(code_lang)
    link = sanitize_link(link)
    if not name:
        return None
    if not link and not (code or '').strip():
        return None
    st = get_current_storage()
    new_id = st.add_record(name, link, desc, category, favorite, tags, code, code_lang)
    if new_id is None:
        return None
    invalidate_cache(full=False, groups=True, categories=True, tags=True, url_index=True)
    log_action('add', new_id, name, {'link': link[:100], 'category': category})
    run_hook('on_add', record_id=new_id, name=name, link=link,
             description=desc, category=category, tags=tags,
             code=code, code_lang=code_lang)
    return new_id


def soft_delete_record(record_id):
    get_current_storage().soft_delete_record(record_id)
    invalidate_cache(full=False, groups=True, categories=True, tags=True, trash=True)


def hard_delete_record(record_id):
    name = get_current_storage().hard_delete_record(record_id)
    invalidate_cache(full=True)
    if name is not None:
        run_hook('on_hard_delete', record_id=record_id, name=name)


def restore_from_trash(record_id):
    ok = get_current_storage().restore_from_trash(record_id)
    if ok:
        invalidate_cache(full=False, groups=True, categories=True, tags=True, trash=True)
    return ok


def update_record(record_id, name, link, desc, category, favorite, tags='',
                  code=None, code_lang=None):
    """Обновить запись. Разрешает очистить ссылку, если у записи есть код."""
    name = _truncate(name, MAX_NAME_LEN).strip()
    desc = _truncate(desc, MAX_DESC_LEN).strip()
    category = _truncate(category, MAX_CATEGORY_LEN).strip()
    tags = _truncate(tags, MAX_TAGS_LEN).strip()
    link = sanitize_link(link)
    if code is not None:
        code = _truncate(code, MAX_CODE_LEN)
    if code_lang is not None:
        code_lang = validate_code_lang(code_lang)
    get_current_storage().update_record(record_id, name, link, desc, category,
                                        favorite, tags,
                                        code=code, code_lang=code_lang)
    invalidate_cache(full=True)


def update_category_only(record_id, category):
    info = get_current_storage().update_category_only(record_id, category)
    if not info:
        return False
    invalidate_cache(full=True)
    run_hook('on_category_change', record_id=record_id, name=info['name'],
             old_category=info['old'], new_category=info['new'])
    return True


def update_description(record_id, desc):
    ok = get_current_storage().update_description(record_id, desc)
    if ok:
        invalidate_cache(full=False, groups=True)
    return ok


def move_record_to_group(record_id, new_group):
    info = get_current_storage().move_record_to_group(record_id, new_group)
    if not info:
        return False
    invalidate_cache(full=True)
    run_hook('on_move_to_group', record_id=record_id,
             from_group=info['old'], to_group=info['new'])
    return True


def toggle_favorite_record(record_id):
    get_current_storage().toggle_favorite_record(record_id)
    invalidate_cache(full=True)


def increment_click(record_id):
    info = get_current_storage().increment_click(record_id)
    if info:
        run_hook('on_click', record_id=record_id, name=info['name'],
                 link=info['link'], click_count=info['count'])


def update_order(group_name, ordered_ids):
    get_current_storage().update_order(group_name, ordered_ids)
    invalidate_cache(full=True)


def sort_links_in_group(group_name, sort_by='name', asc=True):
    ok = get_current_storage().sort_links_in_group(group_name, sort_by, asc)
    if ok:
        invalidate_cache(full=True)
    return ok


def get_record_by_id(record_id):
    return get_current_storage().get_record_by_id(record_id)


def get_all_categories(ws=None):
    return get_current_storage().get_all_categories()


def get_all_tags(ws=None):
    return get_current_storage().get_all_tags()


# ============================================================
# ВЫБОРКИ
# ============================================================
def format_datetime(dt_str):
    if not dt_str:
        return ''
    try:
        return datetime.strptime(str(dt_str), '%Y-%m-%d %H:%M:%S').strftime('%d.%m.%Y %H:%M')
    except Exception:
        return str(dt_str)


def is_network_path(link):
    return link.startswith('\\\\') or link.startswith('//')


def format_link_display(link):
    if not link:
        return '📝'
    if is_network_path(link):
        return '📁'
    try:
        from urllib.parse import urlparse
        p = urlparse(link)
        if p.scheme in ('http', 'https') and p.netloc:
            host = p.netloc.split(':')[0]
            if host.startswith(('192.168.', '10.', '127.', '172.')):
                return '🏠'
            return '🌐'
    except Exception:
        pass
    if len(link) > 1 and link[1] == ':':
        return '💾'
    if link.startswith('mailto:'):
        return '📧'
    if link.startswith('file://'):
        return '🗄️'
    return '🔗'


def group_records(ws, include_deleted=False):
    groups = {}
    for row in ws.iter_rows(min_row=2, values_only=True):
        name = safe_str(row[1]).strip() if row[1] else ''
        if not name:
            continue
        deleted_at = safe_str(row[10]) if len(row) > 10 else ''
        if deleted_at and not include_deleted:
            continue
        groups.setdefault(name, [])
        tags_str = safe_str(row[11]) if len(row) > 11 else ''
        tags = [t.strip() for t in tags_str.split(',') if t.strip()] if tags_str else []
        groups[name].append({
            'id': row[0],
            'link': safe_str(row[2]).strip(),
            'description': safe_str(row[3]).strip(),
            'category': safe_str(row[4]).strip(),
            'favorite': safe_int(row[6] if len(row) > 6 else 0, 0),
            'order_index': safe_int(row[7] if len(row) > 7 else 0, 0),
            'created_at': format_datetime(row[8] if len(row) > 8 else ''),
            'updated_at': format_datetime(row[9] if len(row) > 9 else ''),
            'updated_at_iso': safe_str(row[9]) if len(row) > 9 else '',
            'deleted_at': deleted_at,
            'tags': tags,
            'tags_str': tags_str,
            'click_count': safe_int(row[12] if len(row) > 12 else 0, 0),
            'last_clicked_at': safe_str(row[13]) if len(row) > 13 else '',
            'code': safe_str(row[14]) if len(row) > 14 else '',
            'code_lang': safe_str(row[15]) if len(row) > 15 else '',
            'status': 'unknown',
        })
    for name in groups:
        groups[name].sort(key=lambda x: x['order_index'])
    return groups


def get_names_for_category(ws, category):
    names = set()
    for row in ws.iter_rows(min_row=2, values_only=True):
        if safe_str(row[10]) if len(row) > 10 else '':
            continue
        if row[4] == category and row[1]:
            names.add(safe_str(row[1]).strip())
    return sorted(names)


def get_cached_data():
    """Кэш данных. Сбрасывается через invalidate_cache."""
    global _CACHE
    if _CACHE['groups'] is not None:
        return _CACHE
    wb = read_workbook(); ws = wb.active
    groups = group_records(ws)
    categories = get_all_categories(ws)
    all_tags = get_all_tags(ws)
    total = 0
    category_counts = {}
    all_names = set()
    url_index = {}
    for name, items in groups.items():
        total += len(items)
        all_names.add(name)
        for item in items:
            cat = item['category']
            if cat:
                category_counts[cat] = category_counts.get(cat, 0) + 1
            key = normalize_url_for_dup(item['link'])
            if key:
                url_index.setdefault(key, []).append({
                    'id': item['id'], 'name': name,
                    'category': cat, 'link': item['link']
                })
    uncategorized = total - sum(category_counts.values())
    try:
        trash_count = get_current_storage().get_trash_count()
    except Exception:
        trash_count = 0
    _CACHE.update({
        'groups': groups, 'categories': categories,
        'category_counts': category_counts, 'total': total,
        'all_names': sorted(all_names), 'all_tags': all_tags,
        'trash_count': trash_count,
        'uncategorized_count': max(uncategorized, 0),
        'url_index': url_index,
    })
    return _CACHE


# ============================================================
# HTML-БИЛДЕРЫ
# ============================================================
def _build_sidebar_cats(all_categories, category_counts, category_colors, total, uncategorized_count=0):
    all_btn = (f'<button class="cat-btn active" data-filter="all">'
               f'<span class="cat-indicator"></span>'
               f'<span class="cat-name">Все</span>'
               f'<span class="badge">{total}</span></button>')
    none_btn = ''
    if uncategorized_count > 0:
        none_btn = (f'<button class="cat-btn cat-btn-none" data-filter="__none__" '
                    f'style="--cat-color:#888888;">'
                    f'<span class="cat-indicator"></span>'
                    f'<span class="cat-name">Без категории</span>'
                    f'<span class="badge">{uncategorized_count}</span></button>')
    user_html = ''
    for cat in all_categories:
        count = category_counts.get(cat, 0)
        color = category_colors.get(cat, '')
        if color and not is_valid_hex_color(color):
            color = ''
        style = f' style="--cat-color:{escape(color)};"' if color else ''
        user_html += (f'<button class="cat-btn" data-filter="{escape(cat)}" '
                      f'data-cat-name="{escape(cat)}" data-cat-count="{count}"{style}>'
                      f'<span class="cat-indicator"></span>'
                      f'<span class="cat-name">{escape(cat)}</span>'
                      f'<span class="badge">{count}</span></button>')
    return all_btn + none_btn + f'<div id="userCategoriesList">{user_html}</div>'


def _build_cat_options(all_categories):
    html = '<option value="">— выберите или введите свою —</option>'
    for cat in all_categories:
        html += f'<option value="{escape(cat)}">{escape(cat)}</option>'
    return html + '<option value="__other__">➕ Другая...</option>'


def _build_datalist(groups):
    return ''.join([f'<option value="{escape(n)}">' for n in sorted(set(groups.keys()))])


def _build_cards(groups, category_colors):
    parts = []
    group_icons = load_group_icons()
    for card_idx, (name, items) in enumerate(groups.items()):
        links_parts = []
        card_cats = set()
        for item in items:
            cat = item['category']
            if cat:
                card_cats.add(cat)
            link = item['link']
            desc = item['description']
            category = item['category']
            link_id = item['id']
            favorite = item['favorite']
            created = item['created_at']
            updated = item['updated_at']
            updated_iso = item['updated_at_iso']
            tags = item.get('tags', [])
            click_count = item.get('click_count', 0)
            code_text = item.get('code', '') or ''
            code_lang = item.get('code_lang', '') or ''
            has_link = bool(link)

            link_esc = escape(link)
            if has_link:
                if is_network_path(link):
                    href, target = 'file://' + link.replace('\\', '/'), 'target="_blank"'
                else:
                    href, target = link, 'target="_blank" rel="noopener noreferrer"'
            else:
                href, target = '', ''

            color_style = ''
            if category and category in category_colors:
                cc = category_colors[category]
                if is_valid_hex_color(cc):
                    color_style = f'background: {cc}33; border-color: {cc}; color: {cc};'
            category_tag = (f'<span class="category-tag" style="{color_style}" '
                            f'data-category="{escape(category)}" role="button" tabindex="0" '
                            f'title="Показать только эту категорию">'
                            f'<span class="meta-ico">🏷️</span>{escape(category)}</span>') if category else ''

            fav_class = 'fav-on' if favorite else 'fav-off'
            fav_icon = '⭐' if favorite else '☆'
            fav_html = f'<span class="favorite-btn {fav_class}" data-id="{link_id}" title="Избранное">{fav_icon}</span>'
            item_fav_class = 'is-favorite' if favorite else ''

            if has_link:
                icon_esc = escape(format_link_display(link))
            else:
                icon_esc = '💻'

            tags_str_attr = escape(','.join(tags))
            desc_html = (f'<div class="desc" data-original="{escape(desc)}" data-id="{link_id}" '
                         f'title="Двойной клик — редактировать">{escape(desc)}</div>')

            meta_left_parts = []
            if click_count > 0:
                meta_left_parts.append(f'<span class="meta-item" title="Просмотров: {click_count}">'
                                       f'<span class="meta-ico">👁</span>{click_count}</span>')
            for t in tags:
                meta_left_parts.append(f'<span class="tag-chip" data-tag="{escape(t)}" '
                                       f'role="button" tabindex="0" '
                                       f'title="Показать все ссылки с тегом #{escape(t)}">'
                                       f'#{escape(t)}</span>')
            meta_right_parts = []
            if created:
                meta_right_parts.append(f'<span class="meta-item meta-created" data-id="{link_id}" '
                                        f'role="button" tabindex="0" title="История">'
                                        f'<span class="meta-ico">🕒</span>{escape(created)}</span>')
            if updated and updated != created:
                meta_right_parts.append(f'<span class="meta-item" title="Изменено">'
                                        f'<span class="meta-ico">✏️</span>{escape(updated)}</span>')
            if category_tag:
                meta_right_parts.append(category_tag)
            meta_html = ''
            if meta_left_parts or meta_right_parts:
                meta_html = ('<div class="link-meta">'
                             f'<span class="meta-left">{"".join(meta_left_parts)}</span>'
                             f'<span class="meta-right">{"".join(meta_right_parts)}</span>'
                             '</div>')

            # --- Ссылка (или заглушка «только код») ---
            if has_link:
                link_html = (
                    f'<a href="{escape(href)}" {target} '
                    f'data-full-link="{link_esc}" data-id="{link_id}" '
                    f'class="smart-link" draggable="false" title="Клик — копировать">'
                    f'<span class="link-icon">{icon_esc}</span>'
                    f'<span class="link-display">{escape(link)}</span>'
                    f'</a>'
                )
            else:
                link_html = (
                    f'<span class="smart-link smart-link-code-only" data-id="{link_id}" '
                    f'title="Запись без ссылки — только код">'
                    f'<span class="link-icon">{icon_esc}</span>'
                    f'<span class="link-display">— только код —</span>'
                    f'</span>'
                )

            # --- Блок кода ---
            code_html = ''
            if code_text.strip():
                lang_attr = code_lang or 'plaintext'
                lang_label = (code_lang or 'text').upper()
                code_esc = escape(code_text)
                code_html = (
                    f'<div class="link-code" data-lang="{escape(lang_attr)}" data-id="{link_id}">'
                    f'  <div class="code-header">'
                    f'    <span class="code-lang">{escape(lang_label)}</span>'
                    f'    <span class="code-actions">'
                    f'      <button type="button" class="code-btn code-copy-btn" title="Скопировать код">📋</button>'
                    f'      <button type="button" class="code-btn code-expand-btn" title="Развернуть / свернуть">⤢</button>'
                    f'    </span>'
                    f'  </div>'
                    f'  <pre class="code-body"><code class="language-{escape(lang_attr)}">{code_esc}</code></pre>'
                    f'</div>'
                )

            links_parts.append(f'''
            <div class="link-item {item_fav_class}" data-id="{link_id}" data-category="{escape(category)}" data-favorite="{favorite}" data-updated="{escape(updated_iso)}" data-status="unknown" data-tags="{tags_str_attr}">
                <div class="link-main">
                    <input type="checkbox" class="bulk-checkbox" data-id="{link_id}" aria-label="Выбрать">
                    {fav_html}
                    {link_html}
                    <div class="link-actions">
                        <button class="edit-btn" data-id="{link_id}" title="Редактировать">✏️</button>
                        <button class="delete-btn" data-id="{link_id}" title="Удалить">🗑</button>
                        <span class="link-status" title="Статус"></span>
                    </div>
                </div>
                {desc_html}
                {code_html}
                {meta_html}
            </div>''')

        links_html = ''.join(links_parts)
        card_cats_json = json.dumps(sorted(card_cats), ensure_ascii=False)
        card_color = ''
        for c in card_cats:
            if c in category_colors and is_valid_hex_color(category_colors[c]):
                card_color = category_colors[c]
                break
        color_attr = f' data-cat-color="1" style="--card-cat-color: {card_color};"' if card_color else ''
        group_icon_esc = escape(group_icons.get(name, '📌'))
        name_attr_esc = escape(name)
        parts.append(f'''
        <div class="card" data-categories="{escape(card_cats_json)}" data-group="{escape(name)}" data-updated="{escape(items[0]['updated_at_iso'] if items else '')}"{color_attr}>
            <span class="card-shimmer-frame" aria-hidden="true"></span>
            <div class="card-header">
                <input type="checkbox" class="bulk-checkbox-card" data-group="{escape(name)}" aria-label="Выбрать карточку">
                <span class="card-icon" data-group-name="{name_attr_esc}" title="Сменить иконку">{group_icon_esc}</span>
                <span class="card-title" id="group-name-{card_idx}" data-original="{escape(name)}">{escape(name)}</span>
                <div class="card-sort-menu">
                    <button class="card-sort-btn" data-group="{name_attr_esc}" title="Сортировка">⇅</button>
                </div>
                <button class="rename-group-btn" data-group-name="{name_attr_esc}" title="Переименовать">✏️</button>
                <span class="card-count">{len(items)}</span>
            </div>
            <div class="links" id="collapse_g{card_idx}">{links_html}</div>
        </div>''')
    return ''.join(parts)


# ============================================================
# CORE API
# ============================================================
class CoreAPI:
    BASE_DIR = BASE_DIR
    EXCEL_FILE = EXCEL_FILE
    BACKUP_DIR = BACKUP_DIR
    LOGS_FILE = LOGS_FILE
    SETTINGS_FILE = SETTINGS_FILE
    SECRET_KEY_FILE = SECRET_KEY_FILE

    def read_workbook(self): return read_workbook()
    def atomic_save_workbook(self, wb): return atomic_save_workbook(wb)
    def get_record_by_id(self, rid): return get_record_by_id(rid)
    def add_record(self, **kw): return add_record(**kw)
    def update_record(self, *a, **kw): return update_record(*a, **kw)
    def soft_delete_record(self, rid): return soft_delete_record(rid)
    def hard_delete_record(self, rid): return hard_delete_record(rid)
    def restore_from_trash(self, rid): return restore_from_trash(rid)
    def update_category_only(self, rid, c): return update_category_only(rid, c)
    def update_description(self, rid, d): return update_description(rid, d)
    def move_record_to_group(self, rid, g): return move_record_to_group(rid, g)
    def toggle_favorite_record(self, rid): return toggle_favorite_record(rid)
    def increment_click(self, rid): return increment_click(rid)
    def get_cached_data(self): return get_cached_data()
    def invalidate_cache(self, **kw): return invalidate_cache(**kw)
    def safe_int(self, v, d=0): return safe_int(v, d)
    def safe_str(self, v, d=''): return safe_str(v, d)
    def sanitize_link(self, link): return sanitize_link(link)
    def escape(self, s): return escape(s)
    def normalize_url_for_dup(self, url): return normalize_url_for_dup(url)
    def format_datetime(self, dt): return format_datetime(dt)
    def is_network_path(self, link): return is_network_path(link)
    def is_valid_hex_color(self, c): return is_valid_hex_color(c)
    def is_valid_datetime_str(self, s): return is_valid_datetime_str(s)
    def load_json(self, path, default): return _load_json(path, default)
    def save_json(self, path, data): return _save_json(path, data)
    def log_action(self, *a, **kw): return log_action(*a, **kw)
    def load_settings(self): return load_settings()
    def save_settings(self, s): return save_settings(s)
    def load_group_icons(self): return load_group_icons()
    def save_group_icons(self, i): return save_group_icons(i)
    def load_category_colors(self): return load_category_colors()
    def save_category_colors(self, c): return save_category_colors(c)
    def load_card_order(self): return load_card_order()
    def save_card_order(self, o): return save_card_order(o)
    def load_category_order(self): return load_category_order()
    def save_category_order(self, o): return save_category_order(o)

    def get_storage_type(self):
        return get_current_storage().type

    def import_from_xlsx(self, path):
        return get_current_storage().import_from_xlsx(path)

    def run_hook(self, hook_name, **kw):
        return run_hook(hook_name, **kw)

    def log(self):
        return logging.getLogger('reestr.plugin')


core = CoreAPI()


# ============================================================
# ГЛАВНАЯ
# ============================================================
@app.route('/')
def index():
    data = get_cached_data()
    groups = data['groups'] or {}
    all_cats = data['categories'] or []
    category_counts = data['category_counts'] or {}
    total = data['total'] or 0
    all_tags = data['all_tags'] or {}
    trash_count = data['trash_count'] or 0
    uncategorized_count = data.get('uncategorized_count', 0)
    category_colors = load_category_colors()
    settings = load_settings()
    ui = collect_ui_slots()
    static_version = str(int(time.time()))
    return render_template(
        'index.html',
        static_version=static_version,
        groups=groups, all_categories=all_cats,
        category_colors=category_colors,
        category_counts=category_counts, total=total,
        all_tags=all_tags, trash_count=trash_count,
        settings=settings, ui=ui,
        sidebar_cats_html=_build_sidebar_cats(all_cats, category_counts, category_colors, total, uncategorized_count),
        cards_html=_build_cards(groups, category_colors),
        cat_options=_build_cat_options(all_cats),
        datalist_options=_build_datalist(groups),
        manual_category_order=load_category_order()
    )


# ============================================================
# STORAGE API
# ============================================================
@app.route('/api/storage/info')
def api_storage_info():
    try:
        st = get_current_storage()
    except Exception as e:
        return jsonify({'success': False, 'error': str(e)}), 500

    info = {
        'success': True,
        'type': st.type,
        'db_path': '',
        'sqlite_exists': False,
        'sqlite_size': 0,
        'excel_exists': False,
        'excel_size': 0,
        'total_records': 0,
        'total_in_trash': 0,
    }

    try:
        if Path(EXCEL_FILE).exists():
            info['excel_exists'] = True
            info['excel_size'] = Path(EXCEL_FILE).stat().st_size
    except Exception:
        pass

    try:
        cfg = (load_settings().get('storage') or {}).get('sqlite') or {}
        rel = cfg.get('path') or 'reestr.db'
        p = Path(rel)
        # NEW: не преобразуем абсолютные и UNC-пути
        if not p.is_absolute() and not (str(rel).startswith('\\\\') or str(rel).startswith('//')):
            p = Path(BASE_DIR) / rel
        if p.exists():
            info['sqlite_exists'] = True
            info['sqlite_size'] = p.stat().st_size
        info['db_path'] = str(p)
    except Exception:
        pass

    try:
        records = st.get_all_records()
        info['total_records'] = len([r for r in records if not r.get('deleted_at')])
        info['total_in_trash'] = len([r for r in records if r.get('deleted_at')])
    except Exception:
        pass

    return jsonify(info)


# NEW: проверка пути без переключения хранилища
@app.route('/api/storage/validate_path', methods=['POST'])
def api_storage_validate_path():
    """Проверяет путь для SQLite-базы, не переключая хранилище."""
    data = request.get_json() or {}
    path_str = str(data.get('path', '')).strip()
    ok, err, resolved, is_net = validate_sqlite_path(path_str, BASE_DIR)
    return jsonify({
        'success': ok,
        'error': err,
        'resolved': resolved,
        'is_network': is_net,
    })


# ИЗМЕНЕНО: валидация пути для SQLite + поддержка UNC и абсолютных путей
@app.route('/api/storage/switch', methods=['POST'])
def api_storage_switch():
    data = request.get_json() or {}
    new_type = str(data.get('type', '')).strip().lower()
    if new_type not in ('excel', 'sqlite'):
        return jsonify({'success': False, 'error': 'type: excel или sqlite'}), 400

    if new_type == 'excel':
        try:
            switch_storage('excel')
        except Exception as e:
            log.exception("[STORAGE] Ошибка переключения на Excel")
            return jsonify({'success': False, 'error': str(e)}), 500
        return jsonify({
            'success': True,
            'type': 'excel',
            'message': 'Хранилище переключено на Excel. Перезагрузите страницу.',
        })

    # --- SQLite: валидируем путь ---
    raw_path = str(data.get('sqlite_path', 'reestr.db')).strip() or 'reestr.db'
    ok, err, resolved, is_net = validate_sqlite_path(raw_path, BASE_DIR)
    if not ok:
        return jsonify({'success': False, 'error': err}), 400

    # Сохраняем тот путь, который ввёл пользователь:
    #   - относительный («reestr.db», «data/db.sqlite») — как есть (портативность)
    #   - абсолютный («D:\\data\\reestr.db») — как есть
    #   - UNC («\\\\server\\share\\reestr.db») — как есть
    stored_path = raw_path

    try:
        switch_storage('sqlite', extra={
            'sqlite': {'path': stored_path},
        })
    except RuntimeError as e:
        return jsonify({'success': False, 'error': str(e)}), 500
    except Exception as e:
        log.exception("[STORAGE] Ошибка переключения на SQLite")
        return jsonify({'success': False, 'error': str(e)}), 500

    resp = {
        'success': True,
        'type': 'sqlite',
        'path': stored_path,
        'resolved': resolved,
        'is_network': bool(is_net),
        'message': 'Хранилище переключено на SQLite. Перезагрузите страницу.',
    }
    if is_net:
        resp['warning'] = (
            'База на сетевом ресурсе. Одновременная работа нескольких '
            'пользователей возможна, но медленнее, чем с локальной БД. '
            'WAL-журнал автоматически отключён.'
        )
    return jsonify(resp)


@app.route('/api/storage/import_xlsx', methods=['POST'])
def api_storage_import_xlsx():
    st = get_current_storage()
    if st.type == 'excel':
        return jsonify({
            'success': False,
            'error': 'Импорт имеет смысл только для SQLite. Excel уже работает с этим файлом.'
        }), 400

    if 'file' in request.files:
        f = request.files['file']
        if not f.filename or not f.filename.lower().endswith('.xlsx'):
            return jsonify({'success': False, 'error': 'Ожидается .xlsx'}), 400
        tmp = str(Path(BASE_DIR) / '_import_storage_tmp.xlsx')
        try:
            f.save(tmp)
        except Exception as e:
            return jsonify({'success': False, 'error': f'Сохранение: {e}'}), 500
        try:
            n = st.import_from_xlsx(tmp)
        except Exception as e:
            try: os.remove(tmp)
            except Exception: pass
            return jsonify({'success': False, 'error': f'Импорт: {e}'}), 500
        try: os.remove(tmp)
        except Exception: pass
        invalidate_cache(full=True)
        log_action('storage_import_xlsx', None, '', {'count': n, 'file': f.filename})
        return jsonify({'success': True, 'imported': n, 'source': f.filename})

    data = request.get_json() or {}
    src_path = str(data.get('path', '')).strip()
    if not src_path:
        return jsonify({'success': False, 'error': 'Нет файла или path'}), 400
    if not Path(src_path).exists():
        return jsonify({'success': False, 'error': f'Файл не найден: {src_path}'}), 404
    try:
        n = st.import_from_xlsx(src_path)
    except Exception as e:
        return jsonify({'success': False, 'error': f'Импорт: {e}'}), 500
    invalidate_cache(full=True)
    log_action('storage_import_xlsx', None, '', {'count': n, 'path': src_path})
    return jsonify({'success': True, 'imported': n, 'source': src_path})


@app.route('/api/storage/export_xlsx', methods=['POST'])
def api_storage_export_xlsx():
    import tempfile
    st = get_current_storage()
    try:
        wb = st.read_workbook()
    except Exception as e:
        return jsonify({'success': False, 'error': str(e)}), 500
    tmp_dir = tempfile.gettempdir()
    out = os.path.join(tmp_dir, f'reestr_export_{int(time.time())}.xlsx')
    try:
        wb.save(out)
    except Exception as e:
        return jsonify({'success': False, 'error': f'Сохранение: {e}'}), 500
    log_action('storage_export_xlsx', None, '', {'type': st.type})
    return send_file(
        out,
        as_attachment=True,
        download_name=f'реестр_{datetime.now().strftime("%Y%m%d_%H%M%S")}.xlsx',
        mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )


@app.route('/api/storage/reset', methods=['POST'])
def api_storage_reset():
    st = get_current_storage()
    if st.type != 'sqlite':
        return jsonify({'success': False, 'error': 'Только для SQLite'}), 400
    data = request.get_json() or {}
    confirm = str(data.get('confirm', ''))
    if confirm != 'YES_DELETE_ALL':
        return jsonify({'success': False, 'error': 'Требуется confirm=YES_DELETE_ALL'}), 400
    try:
        with st._lock:
            st._conn.execute('DELETE FROM links')
            st._conn.commit()
    except Exception as e:
        return jsonify({'success': False, 'error': str(e)}), 500
    invalidate_cache(full=True)
    log_action('storage_reset', None, '', {'type': st.type})
    return jsonify({'success': True})


# ============================================================
# ПЛАГИНЫ API
# ============================================================
@app.route('/api/plugins')
def api_plugins():
    return jsonify(get_loaded_plugins())


@app.route('/api/plugins/list')
def api_plugins_list():
    info = get_plugins_info(PLUGINS_DIR, PLUGINS_CONFIG_FILE)
    info['hooks_active'] = get_loaded_plugins().get('hooks', {})
    return jsonify(info)


@app.route('/api/plugins/read')
def api_plugins_read():
    name = request.args.get('name', '').strip()
    err = validate_plugin_name(name)
    if err:
        return jsonify({'success': False, 'error': err}), 400
    code = read_plugin_code(PLUGINS_DIR, name)
    if code is None:
        return jsonify({'success': False, 'error': 'Файл не найден'}), 404
    return jsonify({'success': True, 'code': code, 'name': name})


@app.route('/api/plugins/save', methods=['POST'])
def api_plugins_save():
    data = request.get_json() or {}
    name = data.get('name', '').strip()
    code = data.get('code', '')
    if not isinstance(code, str) or not code:
        return jsonify({'success': False, 'error': 'Пустой код'}), 400
    err = save_plugin_code(PLUGINS_DIR, name, code, create_backup=True)
    if err:
        return jsonify({'success': False, 'error': err}), 400
    return jsonify({'success': True, 'message': 'Сохранено. Нажмите «Перезагрузить плагины».'})


@app.route('/api/plugins/create', methods=['POST'])
def api_plugins_create():
    data = request.get_json() or {}
    name = data.get('name', '').strip()
    code = data.get('code', '')
    err = validate_plugin_name(name)
    if err:
        return jsonify({'success': False, 'error': err}), 400
    path = Path(PLUGINS_DIR) / f'{name}.py'
    if path.exists():
        return jsonify({'success': False, 'error': 'Имя занято'}), 400
    if not isinstance(code, str) or not code.strip():
        code = ('__version__ = "1.0"\n__description__ = ""\n\n\n'
                'def register(hooks, app, config, core):\n    pass\n')
    err = save_plugin_code(PLUGINS_DIR, name, code, create_backup=False)
    if err:
        return jsonify({'success': False, 'error': err}), 400
    return jsonify({'success': True, 'message': 'Создано. Нажмите «Перезагрузить плагины».'})


@app.route('/api/plugins/delete', methods=['POST'])
def api_plugins_delete():
    data = request.get_json() or {}
    err = delete_plugin(PLUGINS_DIR, data.get('name', '').strip())
    if err:
        return jsonify({'success': False, 'error': err}), 400
    return jsonify({'success': True, 'message': 'Удалено'})


@app.route('/api/plugins/toggle', methods=['POST'])
def api_plugins_toggle():
    data = request.get_json() or {}
    err = set_enabled(PLUGINS_CONFIG_FILE, data.get('name', '').strip(),
                      bool(data.get('enabled', True)))
    if err:
        return jsonify({'success': False, 'error': err}), 400
    return jsonify({'success': True})


@app.route('/api/plugins/check', methods=['POST'])
def api_plugins_check():
    data = request.get_json() or {}
    err = check_syntax(data.get('code', ''), f"{data.get('name', 'plugin')}.py")
    if err:
        return jsonify({'success': False, 'error': err})
    return jsonify({'success': True, 'message': 'Синтаксис в порядке'})


@app.route('/api/plugins/reload', methods=['POST'])
def api_plugins_reload():
    try:
        loaded, errors = reload_plugins(
            PLUGINS_DIR, app, PLUGINS_CONFIG_FILE, core=core
        )
        log.info("[RELOAD] Плагинов загружено: %d, ошибок: %d", len(loaded), len(errors))
        return jsonify({
            'success': True,
            'loaded': len(loaded),
            'errors': len(errors),
            'loaded_list': [p['name'] for p in loaded],
            'error_list': errors,
        })
    except Exception as e:
        log.exception("[RELOAD] Ошибка перезагрузки")
        return jsonify({'success': False, 'error': str(e)}), 500


# ============================================================
# КАСТОМНЫЕ ТЕМЫ
# ============================================================
@app.route('/api/custom_themes')
def api_custom_themes():
    return jsonify({'themes': get_custom_themes(CUSTOM_THEMES_FILE)})


@app.route('/api/custom_themes/save', methods=['POST'])
def api_custom_themes_save():
    data = request.get_json() or {}
    err = add_custom_theme(CUSTOM_THEMES_FILE,
                           str(data.get('id', '')).strip().lower(),
                           data.get('theme', {}))
    if err:
        return jsonify({'success': False, 'error': err}), 400
    return jsonify({'success': True})


@app.route('/api/custom_themes/delete', methods=['POST'])
def api_custom_themes_delete():
    data = request.get_json() or {}
    err = delete_custom_theme(CUSTOM_THEMES_FILE, str(data.get('id', '')).strip().lower())
    if err:
        return jsonify({'success': False, 'error': err}), 400
    return jsonify({'success': True})


@app.route('/api/bootstrap')
def api_bootstrap():
    data = get_cached_data()
    return jsonify({
        'group_icons': load_group_icons(),
        'settings': load_settings(),
        'all_tags': data.get('all_tags') or {},
        'categories': data.get('categories') or [],
        'plugins': get_loaded_plugins(),
    })


# ============================================================
# CRUD-РОУТЫ
# ============================================================
@app.route('/add', methods=['POST'])
def add():
    name = request.form.get('name', '').strip()
    link = sanitize_link(request.form.get('link', ''))
    desc = request.form.get('desc', '').strip()
    category = request.form.get('category', '').strip()
    if not category:
        category = request.form.get('category_select', '').strip()
        if category == '__other__':
            category = request.form.get('category_custom', '').strip()
    tags = request.form.get('tags', '').strip()
    group_icon = request.form.get('group_icon', '').strip()[:8]
    code = request.form.get('code', '')
    code_lang = request.form.get('code_lang', '').strip()

    if not name:
        return redirect('/?error=need_name')

    has_code = bool((code or '').strip())
    if not link and not has_code:
        return redirect('/?error=need_link_or_code')

    new_id = add_record(name, link, desc, category, 0, tags, code, code_lang)
    if new_id and group_icon and group_icon != '📌':
        icons = load_group_icons()
        icons[name] = group_icon
        save_group_icons(icons)
    return redirect('/')


@app.route('/delete/<int:rid>', methods=['POST'])
def delete(rid):
    rec = get_record_by_id(rid)
    if not rec:
        return jsonify({'success': False, 'error': 'Не найдена'}), 404
    soft_delete_record(rid)
    log_action('soft_delete', rid, rec['name_del'])
    run_hook('on_delete', record_id=rid, name=rec['name_del'],
             link=rec['link_del'], category=rec['category'])
    return jsonify({'success': True, 'deleted': {
        'name': rec['name_del'], 'link': rec['link_del'],
        'description': rec['description'], 'category': rec['category'],
        'favorite': rec['favorite']}})


@app.route('/restore', methods=['POST'])
def restore():
    data = request.get_json()
    if not data:
        return jsonify({'success': False, 'error': 'Нет данных'}), 400
    name = data.get('name', '').strip()
    link = sanitize_link(data.get('link', ''))
    code = data.get('code', '')
    if not name:
        return jsonify({'success': False, 'error': 'Нет name'}), 400
    if not link and not (code or '').strip():
        return jsonify({'success': False, 'error': 'Нет link или code'}), 400
    new_id = add_record(name=name, link=link, desc=data.get('description', ''),
                        category=data.get('category', ''),
                        favorite=data.get('favorite', 0),
                        code=code, code_lang=data.get('code_lang', ''))
    if new_id is None:
        return jsonify({'success': False, 'error': 'Не удалось'}), 400
    log_action('restore', new_id, name)
    return jsonify({'success': True, 'id': new_id})


@app.route('/get_record/<int:rid>')
def get_record(rid):
    rec = get_record_by_id(rid)
    if not rec:
        return jsonify({'error': 'Не найдена'}), 404
    return jsonify(rec)


@app.route('/edit/<int:rid>', methods=['POST'])
def edit(rid):
    data = request.get_json()
    if not data:
        return 'Ошибка', 400

    rec = get_record_by_id(rid)
    if not rec:
        return 'Запись не найдена', 404

    name = data.get('name', '').strip()
    link = sanitize_link(data.get('link', ''))
    desc = data.get('desc', '').strip()
    category = data.get('category', '').strip()
    tags = data.get('tags', '').strip()

    code = data.get('code', None)
    code_lang = data.get('code_lang', None)
    has_code = bool((code or '').strip())
    if not name:
        return 'Название обязательно', 400
    if not link and not has_code:
        return 'Укажите ссылку или код', 400

    if code is None:
        code = rec.get('code', '') or ''
    if code_lang is None:
        code_lang = rec.get('code_lang', '') or ''

    if not link and not (code or '').strip():
        return 'Нужна хотя бы ссылка или код', 400

    favorite = rec.get('favorite', 0)
    changes = {}
    if rec['name_del'] != name:
        changes['name'] = {'from': rec['name_del'], 'to': name}
    if rec['link_del'] != link:
        changes['link'] = {'from': (rec['link_del'] or '')[:150], 'to': link[:150]}
    if (rec['description'] or '') != desc:
        changes['desc'] = {'from': (rec['description'] or '')[:100], 'to': desc[:100]}
    if (rec['category'] or '') != category:
        changes['category'] = {'from': rec['category'], 'to': category}
    old_code = rec.get('code', '') or ''
    if old_code != (code or ''):
        changes['code'] = {'from': f'(len {len(old_code)})', 'to': f'(len {len(code or "")})'}

    update_record(rid, name, link, desc, category, favorite, tags,
                  code=code, code_lang=code_lang)
    log_action('edit', rid, name, {'changes': changes})
    run_hook('on_edit', record_id=rid, name=name, link=link,
             description=desc, category=category, favorite=favorite, tags=tags,
             code=code, code_lang=code_lang, changes=changes)
    return '', 200


@app.route('/edit_desc/<int:rid>', methods=['POST'])
def edit_desc(rid):
    data = request.get_json() or {}
    desc = safe_str(data.get('desc', '')).strip()
    rec = get_record_by_id(rid)
    old = rec['description'] if rec else ''
    ok = update_description(rid, desc)
    if not ok:
        return jsonify({'success': False, 'error': 'Не найдена'}), 404
    log_action('edit_desc', rid, rec['name_del'] if rec else '',
               {'from': (old or '')[:100], 'to': desc[:100]})
    return jsonify({'success': True, 'desc': desc})


@app.route('/edit_category/<int:rid>', methods=['POST'])
def edit_category(rid):
    data = request.get_json() or {}
    cat = safe_str(data.get('category', '')).strip()
    rec = get_record_by_id(rid)
    old = rec['category'] if rec else ''
    ok = update_category_only(rid, cat)
    if not ok:
        return jsonify({'success': False, 'error': 'Не найдена'}), 404
    log_action('edit_category', rid, rec['name_del'] if rec else '',
               {'from': old, 'to': cat})
    return jsonify({'success': True, 'category': cat})


@app.route('/move_to_group', methods=['POST'])
def move_to_group_route():
    data = request.get_json() or {}
    rid = safe_int(data.get('id', 0), 0)
    group = safe_str(data.get('group_name', '')).strip()
    if not rid or not group:
        return jsonify({'success': False, 'error': 'Неверные данные'}), 400
    rec = get_record_by_id(rid)
    if not rec:
        return jsonify({'success': False, 'error': 'Не найдена'}), 404
    old = rec['name_del']
    if not move_record_to_group(rid, group):
        return jsonify({'success': False, 'error': 'Не удалось'}), 400
    log_action('move_to_group', rid, group, {'from': old, 'to': group})
    return jsonify({'success': True, 'from': old, 'to': group})


@app.route('/toggle_favorite/<int:rid>', methods=['POST'])
def toggle_favorite(rid):
    toggle_favorite_record(rid)
    return '', 200


@app.route('/click/<int:rid>', methods=['POST'])
def click(rid):
    try:
        increment_click(rid)
    except Exception as e:
        log.error("[CLICK] %s", e)
    return '', 204


@app.route('/update_order', methods=['POST'])
def update_order_route():
    data = request.get_json()
    if not data or not data.get('group_name') or not data.get('ordered_ids'):
        return 'Неверные данные', 400
    update_order(data['group_name'], data['ordered_ids'])
    return '', 200


@app.route('/sort_group', methods=['POST'])
def sort_group_route():
    data = request.get_json()
    if not data or not data.get('group_name'):
        return jsonify({'success': False, 'error': 'Нет группы'}), 400
    ok = sort_links_in_group(data['group_name'], data.get('sort_by', 'name'),
                             bool(data.get('asc', True)))
    if ok:
        log_action('sort_group', None, data['group_name'],
                   {'by': data.get('sort_by', 'name'),
                    'asc': bool(data.get('asc', True))})
        return jsonify({'success': True})
    return jsonify({'success': False, 'error': 'Не удалось'}), 400


@app.route('/save_card_order', methods=['POST'])
def save_card_order_route():
    data = request.get_json()
    if not data or not isinstance(data.get('order', []), list):
        return jsonify({'success': False, 'error': 'Неверный формат'}), 400
    save_card_order([str(x) for x in data['order'] if isinstance(x, (str, int))])
    return jsonify({'success': True})


@app.route('/save_category_order', methods=['POST'])
def save_category_order_route():
    data = request.get_json()
    if not data or not isinstance(data.get('order', []), list):
        return jsonify({'success': False, 'error': 'Неверный формат'}), 400
    save_category_order([str(x) for x in data['order'] if isinstance(x, (str, int))])
    return jsonify({'success': True})


@app.route('/update_group_icon', methods=['POST'])
def update_group_icon():
    data = request.get_json() or {}
    group = safe_str(data.get('group', '')).strip()
    icon = safe_str(data.get('icon', '📌')).strip() or '📌'
    if not group:
        return jsonify({'success': False, 'error': 'Нет группы'}), 400
    icons = load_group_icons()
    if icon[:8] == '📌':
        icons.pop(group, None)
    else:
        icons[group] = icon[:8]
    save_group_icons(icons)
    log_action('update_group_icon', None, group, {'icon': icon[:8]})
    return jsonify({'success': True, 'icon': icon[:8]})


@app.route('/get_group_icons')
def get_group_icons():
    return jsonify(load_group_icons())


@app.route('/get_categories')
def get_categories():
    return jsonify(get_cached_data()['categories'] or [])


@app.route('/get_all_tags')
def get_all_tags_route():
    tags = get_cached_data()['all_tags'] or {}
    return jsonify([{'name': k, 'count': v}
                    for k, v in sorted(tags.items(), key=lambda x: -x[1])])


@app.route('/get_names_for_category')
def get_names_for_category_route():
    category = request.args.get('category', '').strip()
    if not category:
        return jsonify([])
    return jsonify(get_names_for_category(read_workbook().active, category))


@app.route('/get_all_names')
def get_all_names():
    return jsonify(get_cached_data()['all_names'] or [])


@app.route('/rename_category', methods=['POST'])
def rename_category():
    data = request.get_json()
    old, new = data.get('old_name', '').strip(), data.get('new_name', '').strip()
    if not old or not new or old == new:
        return jsonify({'success': True})
    new = _truncate(new, MAX_CATEGORY_LEN)
    with _excel_write_lock:
        wb = read_workbook(); ws = wb.active; updated = 0
        for row in ws.iter_rows(min_row=2, max_col=5, values_only=False):
            if row[4].value == old:
                row[4].value = new
                updated += 1
        if updated:
            atomic_save_workbook(wb)
    if not updated:
        return jsonify({'success': False, 'error': 'Не найдена'}), 404
    colors = load_category_colors()
    if old in colors:
        colors[new] = colors.pop(old)
        save_category_colors(colors)
    order = load_category_order()
    if old in order:
        save_category_order([new if x == old else x for x in order])
    invalidate_cache(full=True)
    log_action('rename_category', None, old, {'new_name': new, 'count': updated})
    run_hook('on_category_rename', old_name=old, new_name=new, count=updated)
    return jsonify({'success': True, 'updated': updated})


@app.route('/delete_category', methods=['POST'])
def delete_category():
    data = request.get_json()
    category = data.get('category', '').strip()
    if not category:
        return jsonify({'success': False, 'error': 'Некорректное имя'}), 400
    with _excel_write_lock:
        wb = read_workbook(); ws = wb.active; updated = 0
        for row in ws.iter_rows(min_row=2, max_col=5, values_only=False):
            if row[4].value == category:
                row[4].value = ''
                updated += 1
        if updated:
            atomic_save_workbook(wb)
    if not updated:
        return jsonify({'success': False, 'error': 'Не найдена'}), 404
    colors = load_category_colors()
    if category in colors:
        del colors[category]
        save_category_colors(colors)
    order = load_category_order()
    if category in order:
        save_category_order([x for x in order if x != category])
    invalidate_cache(full=True)
    log_action('delete_category', None, category, {'count': updated})
    run_hook('on_category_delete', category=category, count=updated)
    return jsonify({'success': True, 'updated': updated})


@app.route('/update_category_color', methods=['POST'])
def update_category_color():
    data = request.get_json()
    cat, color = data.get('category', '').strip(), data.get('color', '').strip()
    if not cat or not is_valid_hex_color(color):
        return jsonify({'success': False, 'error': 'Некорректные данные'}), 400
    colors = load_category_colors()
    colors[cat] = color
    save_category_colors(colors)
    return jsonify({'success': True})


@app.route('/rename_group', methods=['POST'])
def rename_group():
    data = request.get_json()
    old, new = data.get('old_name', '').strip(), data.get('new_name', '').strip()
    if not old or not new or old == new:
        return jsonify({'success': True})
    new = _truncate(new, MAX_NAME_LEN)
    with _excel_write_lock:
        wb = read_workbook(); ws = wb.active; updated = 0
        for row in ws.iter_rows(min_row=2, max_col=2, values_only=False):
            if row[1].value == old:
                row[1].value = new
                updated += 1
        if updated:
            atomic_save_workbook(wb)
    if not updated:
        return jsonify({'success': False, 'error': 'Не найдена'}), 404
    order = load_card_order()
    if old in order:
        i = order.index(old); order[i] = new
        save_card_order(order)
    icons = load_group_icons()
    if old in icons:
        icons[new] = icons.pop(old)
        save_group_icons(icons)
    invalidate_cache(full=True)
    log_action('rename_group', None, old, {'new_name': new, 'count': updated})
    run_hook('on_group_rename', old_name=old, new_name=new, count=updated)
    return jsonify({'success': True, 'updated': updated})


# ============================================================
# BULK
# ============================================================
@app.route('/bulk/move', methods=['POST'])
def bulk_move():
    data = request.get_json()
    idset = _parse_ids(data.get('ids', []))
    if idset is None:
        return jsonify({'success': False, 'error': 'Неверные данные'}), 400
    category = _truncate(data.get('category', '').strip(), MAX_CATEGORY_LEN)
    snapshot = []
    with _excel_write_lock:
        wb = read_workbook(); ws = wb.active; count = 0
        now = datetime.now().strftime('%Y-%m-%d %H:%M:%S')
        for idx, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
            rid = safe_int(row[0], -1)
            if rid in idset:
                snapshot.append({'id': rid, 'old_category': safe_str(row[4])})
                ws.cell(row=idx, column=5, value=category)
                ws.cell(row=idx, column=10, value=now)
                count += 1
        atomic_save_workbook(wb)
    invalidate_cache(full=True)
    log_action('bulk_move', None, '', {'count': count, 'category': category})
    run_hook('on_bulk_move', count=count, category=category, ids=list(idset))
    return jsonify({'success': True, 'count': count, 'snapshot': snapshot})


@app.route('/bulk/favorite', methods=['POST'])
def bulk_favorite():
    data = request.get_json()
    idset = _parse_ids(data.get('ids', []))
    if idset is None:
        return jsonify({'success': False, 'error': 'Неверные данные'}), 400
    value = 1 if data.get('value', True) else 0
    snapshot = []
    with _excel_write_lock:
        wb = read_workbook(); ws = wb.active; count = 0
        for idx, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
            rid = safe_int(row[0], -1)
            if rid in idset:
                snapshot.append({'id': rid,
                                 'old_favorite': safe_int(row[6] if len(row) > 6 else 0, 0)})
                ws.cell(row=idx, column=7, value=value)
                count += 1
        atomic_save_workbook(wb)
    invalidate_cache(full=True)
    log_action('bulk_favorite', None, '', {'count': count, 'value': value})
    return jsonify({'success': True, 'count': count, 'snapshot': snapshot})


@app.route('/bulk/tags', methods=['POST'])
def bulk_tags():
    data = request.get_json()
    idset = _parse_ids(data.get('ids', []))
    tag = data.get('tag', '').strip()
    if idset is None or not tag:
        return jsonify({'success': False, 'error': 'Неверные данные'}), 400
    action = data.get('action', 'add')
    tag = _truncate(tag, 100)
    snapshot = []
    with _excel_write_lock:
        wb = read_workbook(); ws = wb.active; count = 0
        for idx, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
            rid = safe_int(row[0], -1)
            if rid in idset:
                cur = safe_str(row[11] if len(row) > 11 else '')
                snapshot.append({'id': rid, 'old_tags': cur})
                tags = [t.strip() for t in cur.split(',') if t.strip()]
                if action == 'add' and tag not in tags:
                    tags.append(tag)
                elif action == 'remove' and tag in tags:
                    tags.remove(tag)
                elif action == 'set':
                    tags = [tag]
                ws.cell(row=idx, column=12,
                        value=_truncate(','.join(tags), MAX_TAGS_LEN))
                count += 1
        atomic_save_workbook(wb)
    invalidate_cache(full=True)
    log_action('bulk_tags', None, '', {'count': count, 'tag': tag, 'action': action})
    return jsonify({'success': True, 'count': count, 'snapshot': snapshot})


@app.route('/bulk/delete', methods=['POST'])
def bulk_delete():
    data = request.get_json()
    idset = _parse_ids(data.get('ids', []))
    if idset is None:
        return jsonify({'success': False, 'error': 'Неверный формат'}), 400
    now = datetime.now().strftime('%Y-%m-%d %H:%M:%S')
    count = 0
    with _excel_write_lock:
        wb = read_workbook(); ws = wb.active
        for idx, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
            rid = safe_int(row[0], -1)
            if rid in idset:
                if safe_str(row[10]) if len(row) > 10 else '':
                    continue
                ws.cell(row=idx, column=11, value=now)
                count += 1
        if count:
            atomic_save_workbook(wb)
    invalidate_cache(full=True)
    log_action('bulk_delete', None, '', {'count': count})
    run_hook('on_bulk_delete', count=count, ids=list(idset))
    return jsonify({'success': True, 'count': count})


@app.route('/bulk/undo', methods=['POST'])
def bulk_undo():
    data = request.get_json() or {}
    op = safe_str(data.get('op', ''))
    snapshot = data.get('snapshot')
    if not isinstance(snapshot, list) or not snapshot:
        return jsonify({'success': False, 'error': 'Нет данных для отката'}), 400
    restored = 0
    with _excel_write_lock:
        wb = read_workbook(); ws = wb.active
        id_to_row = {}
        for idx, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
            id_to_row[safe_int(row[0], -1)] = idx
        for item in snapshot:
            rid = safe_int(item.get('id'), -1)
            row_idx = id_to_row.get(rid)
            if row_idx is None:
                continue
            if op == 'bulk_move':
                ws.cell(row=row_idx, column=5,
                        value=safe_str(item.get('old_category', '')))
                restored += 1
            elif op == 'bulk_favorite':
                ws.cell(row=row_idx, column=7,
                        value=safe_int(item.get('old_favorite', 0), 0))
                restored += 1
            elif op == 'bulk_tags':
                ws.cell(row=row_idx, column=12,
                        value=_truncate(safe_str(item.get('old_tags', '')), MAX_TAGS_LEN))
                restored += 1
        if restored:
            atomic_save_workbook(wb)
    if restored:
        invalidate_cache(full=True)
        log_action('bulk_undo', None, '', {'op': op, 'count': restored})
    return jsonify({'success': True, 'restored': restored})


# ============================================================
# SETTINGS
# ============================================================
@app.route('/get_settings')
def get_settings():
    return jsonify(load_settings())


@app.route('/save_settings', methods=['POST'])
def save_settings_route():
    data = request.get_json()
    if not isinstance(data, dict):
        return jsonify({'success': False, 'error': 'Неверный формат'}), 400
    cur = load_settings()
    if 'custom_toolbar_buttons' in data and isinstance(data['custom_toolbar_buttons'], list):
        cur['custom_toolbar_buttons'] = [c for c in
            (_sanitize_toolbar_btn(b) for b in data['custom_toolbar_buttons'][:20]) if c]
    if 'profile' in data and isinstance(data['profile'], dict):
        cur['profile'] = {
            'name': str(data['profile'].get('name', 'Admin'))[:50],
            'color': str(data['profile'].get('color', '#f5d76e'))[:20]}
    for k in ('shimmer_color_1', 'shimmer_color_2'):
        if k in data:
            cur[k] = str(data.get(k, ''))[:20]
    if 'header_shimmer' in data:
        v = str(data.get('header_shimmer', 'sweep'))
        if v in ('sweep', 'shine', 'breath', 'wave', 'aurora', 'off'):
            cur['header_shimmer'] = v
    if 'card_shimmer' in data:
        cur['card_shimmer'] = bool(data.get('card_shimmer', True))
    if 'shimmer_speed' in data:
        v = str(data.get('shimmer_speed', 'normal'))
        if v in ('slow', 'normal', 'fast'):
            cur['shimmer_speed'] = v
    if 'shimmer_intensity' in data:
        try:
            v = float(data.get('shimmer_intensity', 1.0))
            if 0.1 <= v <= 3.0:
                cur['shimmer_intensity'] = v
        except (ValueError, TypeError):
            pass
    if 'header_anim' in data:
        cur['header_anim'] = bool(data.get('header_anim', True))
    if 'mesh_style' in data:
        v = str(data.get('mesh_style', 'off'))
        if v in ('off', 'float', 'pulse', 'rotate', 'wave', 'orbit', 'flash', 'aurora'):
            cur['mesh_style'] = v
    if 'card_strip' in data:
        v = str(data.get('card_strip', 'left'))
        if v in ('left', 'hover-fill', 'gradient-full', 'gradient-center',
                 'top', 'bottom', 'frame', 'off'):
            cur['card_strip'] = v
    # NEW: разрешаем сохранение темы на сервере (для FOUC-фикса)
    if 'theme' in data:
        t = str(data.get('theme', 'light'))[:64]
        cur['theme'] = t
    save_settings(cur)
    return jsonify({'success': True, 'settings': cur})


# ============================================================
# SHUTDOWN
# ============================================================
_shutdown_done = threading.Event()


def _cleanup():
    if _shutdown_done.is_set():
        return
    _shutdown_done.set()
    log.info("Завершение работы...")
    try:
        for tmp in Path(BASE_DIR).glob('*.tmp'):
            try: tmp.unlink()
            except Exception: pass
        for tmp in Path(BASE_DIR).glob('*.import_tmp'):
            try: tmp.unlink()
            except Exception: pass
    except Exception as e:
        log.error("Ошибка очистки: %s", e)
    try:
        run_hook('on_shutdown')
    except Exception as e:
        log.error("Ошибка on_shutdown: %s", e)
    try:
        if _storage is not None:
            _storage.close()
    except Exception:
        pass
    log.info("Сервер остановлен.")


atexit.register(_cleanup)


def _signal_handler(signum, frame):
    log.info("Получен сигнал %s", signum)
    _cleanup()
    sys.exit(0)


try:
    signal.signal(signal.SIGINT, _signal_handler)
    if hasattr(signal, 'SIGTERM'):
        signal.signal(signal.SIGTERM, _signal_handler)
except Exception:
    pass


# ============================================================
# RUN
# ============================================================
if __name__ == '__main__':
    Path(BASE_DIR / 'templates').mkdir(exist_ok=True)
    Path(BASE_DIR / 'static').mkdir(exist_ok=True)
    Path(BACKUP_DIR).mkdir(exist_ok=True)
    Path(PLUGINS_DIR).mkdir(exist_ok=True)
    Path(BASE_DIR / 'logs').mkdir(exist_ok=True)
    # --- Одноразовая миграция Excel → SQLite ---
    try:
        _auto_migrate_to_sqlite(BASE_DIR)
    except Exception as e:
        log.error("[MIGRATION] Ошибка: %s", e)
    # ------------------------------------------
    try:
        st = get_current_storage()
        log.info("[STORAGE] Инициализировано: type=%s", st.type)
        if st.type == 'sqlite':
            total = len(st.get_all_records())
            log.info("[STORAGE] Записей в БД: %d", total)
    except Exception as e:
        log.error("[STORAGE] Ошибка инициализации: %s", e)
        raise

    loaded, errors = load_plugins(PLUGINS_DIR, app, PLUGINS_CONFIG_FILE, core=core)

    debug_mode = os.environ.get('REESTR_DEBUG') == '1'
    log.info("=" * 60)
    log.info("РЕЕСТР ПОЛЕЗНЫХ ССЫЛОК")
    log.info("=" * 60)
    log.info("Рабочая папка: %s", BASE_DIR)
    log.info("Debug: %s", 'вкл' if debug_mode else 'выкл (REESTR_DEBUG=1)')
    log.info("Разрешённые Origin: %s", ', '.join(ALLOWED_ORIGINS))
    if loaded:
        log.info("Плагинов загружено: %d", len(loaded))
        for p in loaded:
            mark = '✓' if p.get('enabled', True) else '⊘'
            log.info("  %s %s v%s — %s", mark, p['name'], p['version'],
                     p['description'][:60])
    if errors:
        log.warning("Ошибок загрузки: %d", len(errors))
        for e in errors:
            log.warning("  ✗ %s: %s", e['file'], e['error'])
    log.info("=" * 60)

    try:
        run_hook('on_startup')
    except Exception as e:
        log.error("on_startup: %s", e)

    try:
        app.run(debug=debug_mode, host='127.0.0.1', port=5432, use_reloader=debug_mode)
    finally:
        _cleanup()
